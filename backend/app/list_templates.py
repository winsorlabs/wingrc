"""Read a folder of list templates into list definitions -- roadmap L.1.

The import contract is the template convention itself
(docs/PLAN-list-library.md §3), not hand-authored definitions. Forty of the
44 workbooks in Jarrod's library share one header block per sheet:

    row 1   title                       (merged; text sits in column C)
    row 2   CMMC Practice / NIST 800-171: CM.L2-3.4.3 / 3.4.4
    row 3   HOW TO USE THIS TEMPLATE
    row 4   What goes here: <description>
    row 5   Responsible: <role>   |   Review cadence: <cadence>
    row 9   <section title>
    row 10  <column headers>
    row 11+ rows, including [PLACEHOLDER - reason] examples

The unit is a *sheet*: one workbook can hold several lists sharing rows 1-6
(`Hardware and Software Asset Inventory.xlsx` is 3.4.1a Hardware and
3.4.1b Software).

Pure: no DB. Everything here classifies and reports; nothing resolves a
collision silently. `plan_import()` returns what an apply would create;
`list_library.apply_import()` writes it.
"""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl

# NIST SP 800-171 Rev 2 requirement families, by the second number of the
# requirement id (3.<n>.x). Row 2 writes the family prefix on the first id
# only -- "CM.L2-3.4.3 / 3.4.4" -- so later ids take theirs from here.
_FAMILY = {
    1: "AC", 2: "AT", 3: "AU", 4: "CM", 5: "IA", 6: "IR", 7: "MA",
    8: "MP", 9: "PS", 10: "PE", 11: "RA", 12: "CA", 13: "SC", 14: "SI",
}
_PRACTICE_LINE = re.compile(r"CMMC Practice\s*/\s*NIST 800-171\s*:?(.*)", re.I | re.S)
_REQUIREMENT = re.compile(r"(?:\b([A-Z]{2})\.L2-)?\b3\.(\d{1,2})\.(\d{1,2})\b")
_RESPONSIBLE = re.compile(
    r"Responsible:\s*(?P<resp>.*?)\s*(?:\|\s*Review cadence:\s*(?P<cad>.*))?$", re.I | re.S
)

HEADER_ROW = 10
FIRST_DATA_ROW = 11


@dataclass(frozen=True)
class MergeRule:
    """One recorded de-duplication decision: `absorbed` folds into `target`.

    The device-inventory merge (Jarrod, 2026-10-10): 3.4.1a Hardware and
    3.1.1c Authorized Devices are one list tagged with both practices.
    3.1.1c's finer-grained columns are the schema (Make and Model separate,
    OS and BIOS separate); 3.4.1a's columns map onto them explicitly, and
    two with no equivalent are added. 3.4.1a's `Type` is a device class
    (Workstation/Server/Firewall), not 3.1.1c's `Asset Type` (a CMMC
    category), so it lands as `Device Type` -- the one rename that keeps two
    different facts from sharing a name. 3.4.1a's example rows are not
    carried: they are illustrative, and its combined columns would land
    half-mapped."""

    target: tuple[str, str]
    absorbed: tuple[str, str]
    column_map: dict[str, str]
    added_columns: tuple[str, ...]


DEVICE_MERGE = MergeRule(
    target=("AC/3.1.1/Authorized-Entities.xlsx", "3.1.1c Authorized Devices"),
    absorbed=("CM/3.4.1/Hardware and Software Asset Inventory.xlsx", "3.4.1a Hardware"),
    column_map={
        "Asset Name": "Name",
        "Owner / User": "Owner / Primary User",
        "Make / Model": "Make",
        "Serial / Asset Tag": "Serial # or Asset Tag",
        "OS / Firmware": "OS",
        "Location": "Location",
        "Type": "Device Type",
        "Baseline Ref": "Baseline Ref",
    },
    added_columns=("Device Type", "Baseline Ref"),
)
MERGES: tuple[MergeRule, ...] = (DEVICE_MERGE,)

# A collision Jarrod settled (2026-10-10): when two files claim one practice,
# the conforming one wins and the other is reported, never imported.
SUPERSEDED: dict[str, str] = {
    "AC/3.1.20/CMMC_3.1.20_External_Connections.xlsx":
        "AC/3.1.20/External Systems and Connections.xlsx",
}

# The four existing scope-graph views (catalog.py) these sheets correspond
# to. Recorded on the definition as metadata only in L.1 -- the views keep
# working exactly as before; the scope-graph join itself is L.6.
PROJECTION_VIEWS: dict[tuple[str, str], str] = {
    ("AC/3.1.1/Authorized-Entities.xlsx", "3.1.1a Authorized Users"): "3.1.1a-authorized-users",
    ("AC/3.1.1/Authorized-Entities.xlsx", "3.1.1b Auth Processes"): "3.1.1b-auth-processes",
    ("AC/3.1.1/Authorized-Entities.xlsx", "3.1.1c Authorized Devices"): "3.1.1c-authorized-devices",
    ("AC/3.1.1/Authorized-Entities.xlsx", "External Services"): "external-services",
}


@dataclass
class SheetResult:
    file: str  # path relative to the library root, forward slashes
    sheet: str
    status: str  # importable | non_conforming | needs_normalization | merged | superseded
    reason: str = ""
    workbook_title: str = ""
    title: str = ""
    description: str = ""
    responsible: str = ""
    review_cadence: str = ""
    control_keys: tuple[str, ...] = ()
    columns: tuple[str, ...] = ()
    rows: tuple[tuple[str, ...], ...] = ()
    sections: tuple[str, ...] = ()
    file_sha256: str = ""
    changelogs: tuple[str, ...] = ()  # paths relative to the Changelog root
    projection_view_id: str | None = None
    absorbed: tuple[str, ...] = ()  # "file#sheet" of sheets merged into this one

    @property
    def source_ref(self) -> str:
        return f"{self.file}#{self.sheet}"

    @property
    def list_key(self) -> str:
        return slug(self.sheet)


@dataclass
class ImportPlan:
    sheets: list[SheetResult] = field(default_factory=list)
    not_ingested: list[str] = field(default_factory=list)  # Archive/, by decision

    def of(self, *statuses: str) -> list[SheetResult]:
        return [s for s in self.sheets if s.status in statuses]

    @property
    def definitions(self) -> list[SheetResult]:
        """What an apply would create: importable sheets, merges applied."""
        return self.of("importable")


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")


def control_keys(practice_text: str) -> tuple[str, ...]:
    """'CM.L2-3.4.3 / 3.4.4' -> ('CM.L2-3.4.3', 'CM.L2-3.4.4'), in order,
    deduplicated. A bare id takes its family from the requirement number."""
    keys: list[str] = []
    for fam, mid, last in _REQUIREMENT.findall(practice_text):
        family = fam or _FAMILY.get(int(mid), "")
        if not family:
            continue
        key = f"{family}.L2-3.{int(mid)}.{int(last)}"
        if key not in keys:
            keys.append(key)
    return tuple(keys)


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        midnight = value.time() == datetime.min.time()
        return value.date().isoformat() if midnight else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _row_cells(ws, r: int) -> list[tuple[int, str]]:
    return [(c.column, _cell_text(c.value)) for c in ws[r] if _cell_text(c.value)]


def _row_text(ws, r: int) -> str:
    return " ".join(t for _, t in _row_cells(ws, r))


def _wide_repeat(columns: list[str]) -> tuple[str, ...] | None:
    """The shorter column tuple a wide sheet repeats, if it is one --
    'rows wearing a costume' (plan §4.2)."""
    for k in range(1, len(columns) // 2 + 1):
        if len(columns) % k:
            continue
        unit = columns[:k]
        if all(columns[i:i + k] == unit for i in range(0, len(columns), k)):
            return tuple(unit)
    return None


def read_sheet(ws, rel_file: str, file_sha256: str) -> SheetResult:
    res = SheetResult(
        file=rel_file, sheet=ws.title, status="non_conforming", file_sha256=file_sha256
    )
    title = _row_text(ws, 1)
    practice = _PRACTICE_LINE.search(_row_text(ws, 2))
    how_to = _row_text(ws, 3).upper().startswith("HOW TO USE")
    if not (title and practice and how_to):
        res.reason = "does not follow the template header block (rows 1-3)"
        return res

    res.workbook_title = title
    res.control_keys = control_keys(practice.group(1))
    if not res.control_keys:
        res.reason = "row 2 names no NIST 800-171 practice id"
        return res
    desc = _row_text(ws, 4)
    res.description = re.sub(r"^What goes here:\s*", "", desc, flags=re.I)
    m = _RESPONSIBLE.match(_row_text(ws, 5))
    if m:
        res.responsible = (m.group("resp") or "").strip()
        res.review_cadence = (m.group("cad") or "").strip()

    section_cells = _row_cells(ws, 9)
    res.title = section_cells[0][1] if section_cells else title
    header = _row_cells(ws, HEADER_ROW)
    res.columns = tuple(t for _, t in header)
    positions = [col for col, _ in header]
    if not res.columns:
        res.reason = f"row {HEADER_ROW} holds no column headers"
        return res

    if len(section_cells) > 1:
        res.status = "needs_normalization"
        res.sections = tuple(t for _, t in section_cells)
        res.reason = (
            f"wide: {len(section_cells)} side-by-side sections in row 9, each a "
            "column group -- rows wearing a costume; normalize by hand, not guessed here"
        )
        return res
    unit = _wide_repeat(list(res.columns))
    if unit:
        res.status = "needs_normalization"
        res.reason = (
            f"wide: the columns {list(unit)} repeat {len(res.columns) // len(unit)} times"
        )
        return res

    rows: list[tuple[str, ...]] = []
    sections: list[str] = []
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        cells = _row_cells(ws, r)
        if not cells:
            continue
        nxt = [t for _, t in _row_cells(ws, r + 1)]
        if len(cells) == 1 and len(nxt) >= 2 and nxt[0] == res.columns[0]:
            sections.append(cells[0][1])  # a section title before a repeated header
            continue
        values = [t for _, t in cells]
        if values == list(res.columns)[: len(values)]:
            continue  # a repeated header row
        by_col = dict(cells)
        rows.append(tuple(by_col.get(col, "") for col in positions))
    if sections:
        res.status = "needs_normalization"
        res.sections = tuple(sections)
        res.reason = (
            f"multi-section: {len(sections) + 1} sections under one tab header; "
            "the tool/section is a column value, not a list boundary -- not flattened here"
        )
        return res

    res.rows = tuple(rows)
    res.status = "importable"
    res.projection_view_id = PROJECTION_VIEWS.get((rel_file, ws.title))
    return res


def _match_changelogs(changelog_root: Path | None, sheets: list[SheetResult]) -> None:
    if changelog_root is None or not changelog_root.is_dir():
        return
    logs = {
        p.relative_to(changelog_root).as_posix(): p.read_text(encoding="utf-8", errors="replace")
        for p in sorted(changelog_root.rglob("*.md"))
    }
    for s in sheets:
        stem = Path(s.file).stem
        needles = {stem.lower(), s.sheet.lower()}
        s.changelogs = tuple(
            name for name, text in logs.items() if any(n in text.lower() for n in needles)
        )


def plan_import(
    lists_root: Path, changelog_root: Path | None = None, archive_root: Path | None = None
) -> ImportPlan:
    plan = ImportPlan()
    for path in sorted(lists_root.rglob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        rel = path.relative_to(lists_root).as_posix()
        data = path.read_bytes()
        sha = hashlib.sha256(data).hexdigest()
        wb = openpyxl.load_workbook(path, data_only=True)
        sheets = [read_sheet(ws, rel, sha) for ws in wb.worksheets]
        superseded = (
            f"; superseded by {SUPERSEDED[rel]} (decision 2026-10-10)" if rel in SUPERSEDED else ""
        )
        if all(s.status == "non_conforming" for s in sheets):
            plan.sheets.append(SheetResult(
                file=rel, sheet="*", status="non_conforming", file_sha256=sha,
                reason=f"no sheet follows the template header block ({len(sheets)} sheets); "
                "normalize by hand -- no second parser is written for this file" + superseded,
            ))
            continue
        if superseded:
            plan.sheets.append(SheetResult(
                file=rel, sheet="*", status="superseded", file_sha256=sha, reason=superseded[2:],
            ))
            continue
        plan.sheets.extend(sheets)

    by_ref = {(s.file, s.sheet): s for s in plan.sheets}
    for rule in MERGES:
        target, absorbed = by_ref.get(rule.target), by_ref.get(rule.absorbed)
        if not (target and absorbed and target.status == absorbed.status == "importable"):
            continue
        unmapped = [c for c in absorbed.columns if c not in rule.column_map]
        if unmapped:
            raise ValueError(
                f"merge {absorbed.sheet} -> {target.sheet}: no recorded mapping for {unmapped}; "
                "a merge is never guessed"
            )
        target.columns = target.columns + tuple(
            c for c in rule.added_columns if c not in target.columns
        )
        target.rows = tuple(r + ("",) * (len(target.columns) - len(r)) for r in target.rows)
        target.control_keys = tuple(dict.fromkeys(target.control_keys + absorbed.control_keys))
        target.absorbed = target.absorbed + (absorbed.source_ref,)
        absorbed.status = "merged"
        absorbed.reason = f"merged into {target.sheet} (decision 2026-10-10)"

    _match_changelogs(changelog_root, plan.sheets)
    if archive_root is not None and archive_root.is_dir():
        plan.not_ingested = sorted(
            p.relative_to(archive_root.parent).as_posix() for p in archive_root.rglob("*.xlsx")
        )
    return plan
