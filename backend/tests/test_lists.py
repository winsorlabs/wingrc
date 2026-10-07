"""Lists slice (2026-10-07): the on-screen list views, the operator
overlay, the `[PLACEHOLDER - reason]` convention, and the `baseline`
document type. The export's tenant-isolation and read-only-role fixes are
tests/test_list_export.py.

The invariant most of this file is about: an import or sync may not clear
or overwrite an operator-entered field it has no source for. repo.upsert()
used to assign every field wholesale, so applying a Liongard CHANGED row
erased a hand-entered Location -- and, because every fresh pull carries
in_boundary=True, put a device a reviewer had rejected back in scope.
"""
from __future__ import annotations

import uuid
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient
from pydantic import ValidationError
from sqlalchemy import select

from app.auth import get_current_user
from app.catalog import (
    ALL_VIEWS,
    AUTHORIZED_DEVICES,
    AUTHORIZED_PROCESSES,
    AUTHORIZED_USERS,
)
from app.db import get_session
from app.domain import (
    CanonicalEntity,
    ChangeType,
    EntityType,
    ScopeCategory,
    Source,
    placeholder_reason,
)
from app.importers.liongard import device_profile_to_canonical
from app.importers.workbook import parse_workbook
from app.list_projection import project
from app.main import app
from app.models import (
    AssessmentObjective,
    Control,
    ControlState,
    Framework,
    Organization,
    ScopeEntity,
)
from app.reconcile import reconcile
from app.render import render_view_bytes
from app.routers.scope import OperatorOverlayAttributes, ScopeEntityPatch
from tests.conftest import _app_session, _authed, _grant, _make_fake_user

_BIOS_PLACEHOLDER = "[PLACEHOLDER - BIOS FW version not collected]"


def _sheet(xlsx: bytes):
    return openpyxl.load_workbook(BytesIO(xlsx)).active


def _liongard_device(hostname: str = "LT-0001") -> CanonicalEntity:
    entity, _ = device_profile_to_canonical(
        {
            "ID": f"id-{hostname}",
            "EnvironmentID": 8815,
            "InventoryState": "Inventory",
            "Hostname": hostname,
            "SerialNumber": f"SN-{hostname}",
            "MACAddress": ["14:9d:99:8b:72:36"],
            "Manufacturer": "Dell",
            "Model": "Latitude 7440",
            "OperatingSystem": "Windows 11 Pro",
            "Type": "laptop",
        },
        source_ref="liongard:8815",
    )
    assert entity is not None
    return entity


# ---------------------------------------------------------------------------
# No database
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("view", ALL_VIEWS, ids=lambda v: v.id)
def test_rendered_header_row_matches_the_view_columns_in_order(view):
    """Static guard: a column added to the catalog cannot silently skip the
    sheet, or the reverse."""
    ws = _sheet(render_view_bytes(view, []))
    header = [c.value for c in ws[5] if c.value is not None]
    assert header == [display for _, display in view.columns]
    assert ws.title == view.sheet_title


def test_authorized_devices_header_is_the_workbook_header():
    ws = _sheet(render_view_bytes(AUTHORIZED_DEVICES, []))
    assert [c.value for c in ws[5]][:13] == [
        "Name", "Owner / Primary User", "Make", "Model", "Device Subtype",
        "Serial # or Asset Tag", "Mac Address", "OS", "BIOS FW Ver", "Location",
        "Asset Type", "In Service Date", "Decommissioned Date",
    ]


def test_placeholder_reason_parses_only_the_reason_form():
    assert placeholder_reason(_BIOS_PLACEHOLDER) == "BIOS FW version not collected"
    assert placeholder_reason("[placeholder - not in Datto RMM export]") == (
        "not in Datto RMM export"
    )
    assert placeholder_reason("[placeholder]") is None
    assert placeholder_reason("[PLACEHOLDER - ]") is None
    assert placeholder_reason("1.20.0") is None
    assert placeholder_reason(None) is None


def test_placeholder_cell_round_trips_with_its_reason(tmp_path: Path):
    device = CanonicalEntity(
        entity_type=EntityType.DEVICE,
        natural_key="ASSET-0001",
        attributes={
            "Name": "WS-0001",
            "Serial # or Asset Tag": "ASSET-0001",
            "BIOS FW Ver": _BIOS_PLACEHOLDER,
            "location": "[PLACEHOLDER - not in Datto RMM export]",
        },
    )
    [row] = project(AUTHORIZED_DEVICES, [device])
    cols = [k for k, _ in AUTHORIZED_DEVICES.columns]
    bios = row.cells[cols.index("BIOS FW Ver")]
    assert bios.value == _BIOS_PLACEHOLDER
    assert bios.placeholder_reason == "BIOS FW version not collected"
    assert row.cells[cols.index("OS")].value == ""
    assert row.cells[cols.index("OS")].placeholder_reason is None

    xlsx = render_view_bytes(AUTHORIZED_DEVICES, [device])
    ws = _sheet(xlsx)
    cell = ws.cell(row=6, column=cols.index("BIOS FW Ver") + 1)
    assert cell.value == _BIOS_PLACEHOLDER
    assert cell.font.italic, "a placeholder must render distinctly from an empty cell"
    assert ws.cell(row=6, column=cols.index("OS") + 1).value in (None, "")

    # The rendered sheet read back by the workbook importer keeps the
    # reason verbatim -- the reason form is not the illustrative-row token
    # the importer skips.
    path = tmp_path / "devices.xlsx"
    path.write_bytes(xlsx)
    [parsed] = parse_workbook(path)
    assert parsed.attributes["BIOS FW Ver"] == _BIOS_PLACEHOLDER
    assert parsed.attributes["Location"] == "[PLACEHOLDER - not in Datto RMM export]"


def test_liongard_device_fills_the_workbook_columns():
    """Liongard writes canonical keys, not workbook headers; before column
    sources existed a synced org's 3.1.1c rendered Name/Make/Model/OS blank."""
    entity = _liongard_device("LT-0042")
    entity.scope_category = ScopeCategory.CUI_ASSET
    [row] = project(AUTHORIZED_DEVICES, [entity])
    by_col = dict(zip([d for _, d in AUTHORIZED_DEVICES.columns], row.cells, strict=True))
    assert by_col["Name"].value == "LT-0042"
    assert by_col["Make"].value == "Dell"
    assert by_col["Model"].value == "Latitude 7440"
    assert by_col["Serial # or Asset Tag"].value == "SN-LT-0042"
    assert by_col["Mac Address"].value == "14:9d:99:8b:72:36"
    assert by_col["Asset Type"].value == "CUI Asset"


def test_overlay_wins_over_an_imported_raw_value():
    entity = CanonicalEntity(
        entity_type=EntityType.DEVICE,
        natural_key="A-1",
        attributes={"Location": "HQ - Suite 200", "location": "HQ - Suite 300"},
    )
    [row] = project(AUTHORIZED_DEVICES, [entity])
    cols = [k for k, _ in AUTHORIZED_DEVICES.columns]
    assert row.cells[cols.index("Location")].value == "HQ - Suite 300"


def test_empty_auth_processes_sheet_says_why():
    assert AUTHORIZED_PROCESSES.entity_type == EntityType.PROCESS
    ws = _sheet(render_view_bytes(AUTHORIZED_PROCESSES, []))
    assert "No processes recorded" in ws["A3"].value
    assert "add them manually" in ws["A3"].value
    # The explanation sits above the header, never as a data row a
    # re-import would read as a process.
    assert all(c.value is None for c in ws[6])


def test_a_populated_auth_processes_sheet_has_no_explanation():
    proc = CanonicalEntity(
        entity_type=EntityType.PROCESS,
        natural_key="Liongard inspector agent",
        attributes={"Process Name": "Liongard inspector agent", "Running On": "SRV-01"},
    )
    ws = _sheet(render_view_bytes(AUTHORIZED_PROCESSES, [proc]))
    assert "No processes recorded" not in ws["A3"].value
    assert ws.cell(row=6, column=1).value == "Liongard inspector agent"


def test_reconcile_does_not_report_an_overlay_field_a_sync_never_carries():
    current = _liongard_device()
    current.attributes["location"] = "HQ - Suite 200"
    current.attributes["in_service_date"] = "2025-01-10"
    incoming = _liongard_device()
    [change] = reconcile([current], [incoming]).changes
    assert change.change_type == ChangeType.UNCHANGED


def test_overlay_dates_must_be_iso_or_a_placeholder():
    ok = OperatorOverlayAttributes(
        in_service_date="2025-01-10",
        decommissioned_date="[PLACEHOLDER - disposal record not located]",
    )
    assert ok.in_service_date == "2025-01-10"
    with pytest.raises(ValidationError):
        OperatorOverlayAttributes(in_service_date="Jan 10th")


def test_asset_type_is_constrained_to_the_cmmc_categories():
    with pytest.raises(ValidationError):
        ScopeEntityPatch(scope_category="Laptop")
    for value in ("CUI Asset", "CRMA", "SPA"):
        assert ScopeEntityPatch(scope_category=value).scope_category == value


# ---------------------------------------------------------------------------
# Real HTTP, real RLS
# ---------------------------------------------------------------------------


@pytest.fixture(autouse=True)
def _clear_overrides():
    yield
    app.dependency_overrides.clear()


def _org(db_session) -> uuid.UUID:
    org = Organization(name=f"ListsOrg-{uuid.uuid4().hex[:8]}")
    db_session.add(org)
    db_session.flush()
    return org.id


def _client(db_session, org_id: uuid.UUID, role: str = "msp_engineer") -> TestClient:
    user = _make_fake_user(role=role, org_id=org_id, email=f"{uuid.uuid4().hex[:8]}@example.com")
    _grant(db_session, user)
    app.dependency_overrides[get_session] = _app_session(db_session)
    app.dependency_overrides[get_current_user] = _authed(db_session, user)
    return TestClient(app)


def _apply(client: TestClient, org_id: uuid.UUID, entity: CanonicalEntity, change_type: str):
    r = client.post(
        f"/orgs/{org_id}/imports/workbook/apply",
        json={
            "changes": [
                {
                    "change_type": change_type,
                    "entity_type": entity.entity_type.value,
                    "natural_key": entity.natural_key,
                    "incoming": {
                        "scope_category": (
                            entity.scope_category.value if entity.scope_category else None
                        ),
                        "status": entity.status.value,
                        "in_boundary": entity.in_boundary,
                        "source": entity.source.value,
                        "source_ref": entity.source_ref,
                        "attributes": entity.attributes,
                    },
                }
            ]
        },
    )
    assert r.status_code == 200, r.text


def _row(db_session, org_id: uuid.UUID, natural_key: str) -> ScopeEntity:
    db_session.expire_all()
    return db_session.scalars(
        select(ScopeEntity).where(
            ScopeEntity.org_id == org_id, ScopeEntity.natural_key == natural_key
        )
    ).one()


def _synced_device_with_overlay(db_session, client, org_id) -> ScopeEntity:
    device = _liongard_device()
    _apply(client, org_id, device, "new")
    row = _row(db_session, org_id, device.natural_key)
    r = client.patch(
        f"/orgs/{org_id}/scope/{row.id}",
        json={
            "scope_category": "CUI Asset",
            "attributes": {
                "location": "HQ - Suite 200",
                "in_service_date": "2025-01-10",
                "decommissioned_date": "[PLACEHOLDER - still in service]",
                "requested_by": "Security Officer",
            },
        },
    )
    assert r.status_code == 200, r.text
    return row


@pytest.mark.integration
@pytest.mark.parametrize(
    ("field", "expected"),
    [
        ("location", "HQ - Suite 200"),
        ("in_service_date", "2025-01-10"),
        ("decommissioned_date", "[PLACEHOLDER - still in service]"),
        ("requested_by", "Security Officer"),
    ],
)
def test_overlay_field_survives_a_liongard_sync_that_does_not_mention_it(
    db_session, field, expected
):
    org_id = _org(db_session)
    client = _client(db_session, org_id)
    _synced_device_with_overlay(db_session, client, org_id)

    renamed = _liongard_device()
    renamed.attributes["display_name"] = "LT-0001-RENAMED"
    renamed.attributes["Hostname"] = "LT-0001-RENAMED"
    _apply(client, org_id, renamed, "changed")

    row = _row(db_session, org_id, renamed.natural_key)
    assert row.attributes["display_name"] == "LT-0001-RENAMED", "the sync itself must still land"
    assert row.attributes[field] == expected


@pytest.mark.integration
def test_asset_type_survives_a_liongard_sync(db_session):
    org_id = _org(db_session)
    client = _client(db_session, org_id)
    _synced_device_with_overlay(db_session, client, org_id)
    _apply(client, org_id, _liongard_device(), "changed")
    assert _row(db_session, org_id, "SN-LT-0001").scope_category == "CUI Asset"


@pytest.mark.integration
def test_a_liongard_sync_cannot_set_an_overlay_field(db_session):
    org_id = _org(db_session)
    client = _client(db_session, org_id)
    _synced_device_with_overlay(db_session, client, org_id)
    device = _liongard_device()
    device.attributes["location"] = "somewhere a collector made up"
    _apply(client, org_id, device, "changed")
    assert _row(db_session, org_id, device.natural_key).attributes["location"] == "HQ - Suite 200"


@pytest.mark.integration
def test_a_liongard_sync_does_not_put_a_rejected_device_back_in_scope(db_session):
    org_id = _org(db_session)
    client = _client(db_session, org_id)
    row = _synced_device_with_overlay(db_session, client, org_id)
    assert client.patch(
        f"/orgs/{org_id}/scope/{row.id}", json={"in_boundary": False, "status": "decommissioned"}
    ).status_code == 200

    _apply(client, org_id, _liongard_device(), "changed")  # in_boundary=True, status=active
    row = _row(db_session, org_id, "SN-LT-0001")
    assert row.in_boundary is False
    assert row.status == "decommissioned"


@pytest.mark.integration
def test_a_workbook_import_without_a_location_column_keeps_location(db_session):
    org_id = _org(db_session)
    client = _client(db_session, org_id)
    workbook_device = CanonicalEntity(
        entity_type=EntityType.DEVICE,
        natural_key="ASSET-0009",
        attributes={"Name": "WS-0009", "Serial # or Asset Tag": "ASSET-0009"},
        source=Source.WORKBOOK,
        source_ref="authorized-entities.xlsx",
    )
    _apply(client, org_id, workbook_device, "new")
    row = _row(db_session, org_id, "ASSET-0009")
    client.patch(f"/orgs/{org_id}/scope/{row.id}", json={"attributes": {"location": "Annex"}})

    workbook_device.attributes["Make"] = "Dell"
    _apply(client, org_id, workbook_device, "changed")
    row = _row(db_session, org_id, "ASSET-0009")
    assert row.attributes["Make"] == "Dell"
    assert row.attributes["location"] == "Annex"


@pytest.mark.integration
def test_an_operator_can_still_clear_an_overlay_field(db_session):
    org_id = _org(db_session)
    client = _client(db_session, org_id)
    row = _synced_device_with_overlay(db_session, client, org_id)
    r = client.patch(f"/orgs/{org_id}/scope/{row.id}", json={"attributes": {"location": None}})
    assert r.status_code == 200, r.text
    assert _row(db_session, org_id, "SN-LT-0001").attributes["location"] is None


@pytest.mark.integration
def test_overlay_validation_over_http(db_session):
    org_id = _org(db_session)
    client = _client(db_session, org_id)
    row = _synced_device_with_overlay(db_session, client, org_id)
    assert client.patch(
        f"/orgs/{org_id}/scope/{row.id}", json={"attributes": {"in_service_date": "soon"}}
    ).status_code == 422
    assert client.patch(
        f"/orgs/{org_id}/scope/{row.id}", json={"scope_category": "Workstation"}
    ).status_code == 422


@pytest.mark.integration
def test_list_view_over_http_shows_overlay_and_placeholder(db_session):
    org_id = _org(db_session)
    client = _client(db_session, org_id)
    _synced_device_with_overlay(db_session, client, org_id)

    body = client.get(f"/orgs/{org_id}/lists/{AUTHORIZED_DEVICES.id}").json()
    assert body["columns"] == [d for _, d in AUTHORIZED_DEVICES.columns]
    [row] = body["rows"]
    cells = dict(zip(body["columns"], row["cells"], strict=True))
    assert cells["Location"]["value"] == "HQ - Suite 200"
    assert cells["Asset Type"]["value"] == "CUI Asset"
    assert cells["Decommissioned Date"]["placeholder_reason"] == "still in service"
    assert cells["BIOS FW Ver"] == {"value": "", "placeholder_reason": None}


@pytest.mark.integration
def test_empty_process_list_explains_and_a_manual_process_appears(db_session):
    org_id = _org(db_session)
    client = _client(db_session, org_id)

    body = client.get(f"/orgs/{org_id}/lists/{AUTHORIZED_PROCESSES.id}").json()
    assert body["rows"] == []
    assert "No processes recorded" in body["empty_explanation"]

    r = client.post(
        f"/orgs/{org_id}/scope",
        json={
            "entity_type": "process",
            "natural_key": "Liongard inspector agent",
            "attributes": {
                "Process Name": "Liongard inspector agent",
                "Running On": "SRV-UTIL01",
                "Associated Account": "svc-liongard",
                "Description / Purpose": "Collects configuration for continuous monitoring",
            },
        },
    )
    assert r.status_code == 201, r.text
    body = client.get(f"/orgs/{org_id}/lists/{AUTHORIZED_PROCESSES.id}").json()
    assert [c["value"] for c in body["rows"][0]["cells"]] == [
        "Liongard inspector agent", "SRV-UTIL01", "svc-liongard",
        "Collects configuration for continuous monitoring",
    ]
    summary = {v["id"]: v for v in client.get(f"/orgs/{org_id}/lists").json()}
    assert summary[AUTHORIZED_PROCESSES.id]["row_count"] == 1
    assert summary[AUTHORIZED_USERS.id]["row_count"] == 0


@pytest.mark.integration
def test_list_endpoints_are_tenant_isolated(db_session):
    org_a, org_b = _org(db_session), _org(db_session)
    client_a = _client(db_session, org_a)
    _apply(client_a, org_a, _liongard_device("ALPHA-01"), "new")
    client_b = _client(db_session, org_b)
    _apply(client_b, org_b, _liongard_device("BRAVO-01"), "new")

    rows_b = client_b.get(f"/orgs/{org_b}/lists/{AUTHORIZED_DEVICES.id}").json()["rows"]
    assert [r["natural_key"] for r in rows_b] == ["SN-BRAVO-01"]
    counts_b = {v["id"]: v["row_count"] for v in client_b.get(f"/orgs/{org_b}/lists").json()}
    assert counts_b[AUTHORIZED_DEVICES.id] == 1

    assert client_b.get(f"/orgs/{org_a}/lists").status_code == 403
    assert client_b.get(f"/orgs/{org_a}/lists/{AUTHORIZED_DEVICES.id}").status_code == 403


@pytest.mark.integration
def test_c3pao_assessor_can_read_every_list(db_session):
    org_id = _org(db_session)
    client = _client(db_session, org_id, role="c3pao_assessor")
    assert client.get(f"/orgs/{org_id}/lists").status_code == 200
    for view in ALL_VIEWS:
        assert client.get(f"/orgs/{org_id}/lists/{view.id}").status_code == 200


@pytest.mark.integration
def test_lists_overlay_and_sync_never_touch_sprs_or_control_state(db_session):
    from app.engine import start_assessment

    org_id = _org(db_session)
    fw = Framework(key=f"fw-lists-{uuid.uuid4().hex[:6]}", name="Lists FW", version="r2")
    db_session.add(fw)
    db_session.flush()
    ctrl = Control(
        framework_id=fw.id, control_id="AC.L2-3.1.1", family="AC", title="Authorized access",
        requirement_text="Limit system access.", sprs_weight=5, sequence_order=1,
    )
    db_session.add(ctrl)
    db_session.flush()
    db_session.add(AssessmentObjective(control_id=ctrl.id, objective_key="c", text="Devices"))
    db_session.flush()
    assessment = start_assessment(db_session, org_id, fw.id, "Lists SPRS isolation")
    db_session.commit()
    before_score = assessment.sprs_score

    def _states():
        db_session.expire_all()
        return sorted(
            (cs.id, cs.status, cs.responsibility)
            for cs in db_session.scalars(
                select(ControlState).where(ControlState.assessment_id == assessment.id)
            )
        )

    before_states = _states()
    assert before_states, "the scenario must have a control state to compare"

    client = _client(db_session, org_id)
    _synced_device_with_overlay(db_session, client, org_id)
    _apply(client, org_id, _liongard_device(), "changed")
    for view in ALL_VIEWS:
        assert client.get(f"/orgs/{org_id}/lists/{view.id}").status_code == 200
        assert client.get(f"/orgs/{org_id}/exports/{view.id}").status_code == 200

    db_session.expire_all()
    assert db_session.get(type(assessment), assessment.id).sprs_score == before_score
    assert _states() == before_states


@pytest.mark.integration
def test_baseline_is_a_document_type(db_session):
    org_id = _org(db_session)
    client = _client(db_session, org_id)
    r = client.post(
        f"/orgs/{org_id}/documents",
        json={
            "doc_id": "CM-BL-001", "doc_type": "baseline",
            "title": "Windows 11 Workstation Baseline", "body": "# Scope\n\nAll workstations.",
        },
    )
    assert r.status_code == 201, r.text
    listed = client.get(f"/orgs/{org_id}/documents", params={"doc_type": "baseline"}).json()
    assert [d["doc_id"] for d in listed] == ["CM-BL-001"]
