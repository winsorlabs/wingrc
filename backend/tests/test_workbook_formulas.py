"""No workbook WinGRC emits evaluates what someone typed (2026-10-14).

The cells are other people's words -- a Notes field, a person's name, a
template row -- and the workbooks go to assessors. A string starting with
any of OWASP's formula prefixes (=, +, -, @, tab, carriage return) must
reach the file as text with Excel's quote prefix: shown as typed, never
evaluated, even when the cell is edited. `render._save()` does it for
every workbook; these tests cover both renderers and every scope-drawn
list, and pin that nothing saves a workbook any other way.
"""

from __future__ import annotations

import io
import re
import zipfile
from datetime import UTC, datetime
from pathlib import Path

import openpyxl
import pytest

from app.catalog import ALL_VIEWS
from app.domain import CanonicalEntity
from app.list_library import ListVersionSnap
from app.render import render_list_version_bytes, render_view_bytes

APP = Path(__file__).resolve().parents[1] / "app"

PAYLOADS = (
    "=HYPERLINK(\"http://x\",\"y\")",
    "+cmd|' /C calc'!A0",
    "-2+3+cmd|' /C calc'!A0",
    "@SUM(1+1)*cmd|' /C calc'!A0",
    "\t=1+1",
    "\r=1+1",
)


def _check(xlsx: bytes) -> set[str]:
    """Every prefixed cell is text with the quote prefix, and the file
    holds no formula at all. Returns the payloads found, so a test can
    show it was not vacuous."""
    zf = zipfile.ZipFile(io.BytesIO(xlsx))
    for name in zf.namelist():
        if name.startswith("xl/worksheets/"):
            assert not re.search(rb"<f[ >]", zf.read(name)), f"a formula in {name}"
    wb = openpyxl.load_workbook(io.BytesIO(xlsx))
    found = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v[:1] in {"=", "+", "-", "@", "\t", "\r"}:
                    assert cell.data_type == "s", (cell.coordinate, v)
                    assert cell.quotePrefix, (cell.coordinate, v)
                    found.add(v)
    return found


@pytest.mark.parametrize("view", ALL_VIEWS, ids=lambda v: v.id)
def test_a_scope_list_never_evaluates_a_cell(view):
    """The four scope-drawn lists (render_view_bytes), with a payload in
    every attribute the view reads -- the Notes field, a name."""
    entities = [
        CanonicalEntity(
            entity_type=view.entity_type,
            natural_key=f"K-{i}",
            attributes={key: p for key, _ in view.columns}
            | {display: p for _, display in view.columns},
        )
        for i, p in enumerate(PAYLOADS)
    ]
    found = _check(render_view_bytes(view, entities))
    assert found == set(PAYLOADS), f"{view.id} rendered {found}"


def test_a_list_version_never_evaluates_a_cell():
    """The list library's renderer (the standalone export and the
    bundle's lists/), including the title and the column headers."""
    snap = ListVersionSnap(
        list_id="l", list_key="k", is_template=False, control_keys=("AC.L2-3.1.1",),
        version_id="v", version_number=1, title=PAYLOADS[1], description="",
        responsible="", review_cadence="", columns=("Name", PAYLOADS[3]),
        rows=tuple((p, "plain") for p in PAYLOADS),
        saved_at=datetime(2026, 10, 14, tzinfo=UTC), saved_by=None,
    )
    assert _check(render_list_version_bytes(snap)) == set(PAYLOADS)


def test_an_ordinary_cell_is_left_alone():
    snap = ListVersionSnap(
        list_id="l", list_key="k", is_template=False, control_keys=("AC.L2-3.1.1",),
        version_id="v", version_number=1, title="Devices", description="",
        responsible="", review_cadence="", columns=("Name",), rows=(("WS-01",), ("a-b",)),
        saved_at=datetime(2026, 10, 14, tzinfo=UTC), saved_by=None,
    )
    ws = openpyxl.load_workbook(io.BytesIO(render_list_version_bytes(snap))).active
    assert [ws.cell(row=r, column=1).quotePrefix for r in (7, 8)] == [False, False]


def test_every_workbook_leaves_through_one_save():
    """A renderer that called wb.save() itself would skip the property.
    render.py saves in exactly one place, and nothing else in the app
    builds a workbook."""
    render = (APP / "render.py").read_text(encoding="utf-8")
    start = render.index("def _save(")
    end = render.index("\ndef ", start)
    [save] = [m.start() for m in re.finditer(r"\.save\(", render)]
    assert start < save < end
    builders = sorted(
        p.relative_to(APP).as_posix()
        for p in APP.rglob("*.py")
        if "Workbook(" in p.read_text(encoding="utf-8")
    )
    assert builders == ["render.py"]
