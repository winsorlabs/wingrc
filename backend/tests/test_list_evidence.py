"""List library L.4: a populated list becomes evidence for the practices it
is tagged with (list_evidence.py; docs/PLAN-list-library.md section 5).

What these pin:

- A list is offered only when `list_library.is_populated` says so, and an
  unoffered list says why -- an adopted copy starts unfilled by design.
- Offered, never attached: linking changes no control_state and no SPRS.
- A link names a list version; editing the list leaves the link there, the
  drift is said in all three places, and moving forward writes a new link
  while the old one survives.
- Links are confined to the tag set, and a tag a live link rests on cannot
  be removed -- by a person or by an import -- with the link named.
- A product's deactivation leaves a list link alone, so the tag stays
  blocked.
- The bundle embeds the linked version, logs and hashes it, and two
  bundles of unchanged data match.
- c3pao_assessor sees links and changes none; RLS keeps orgs apart.
- The export carries What goes here, Responsible and Review cadence (T).
"""

from __future__ import annotations

import io
import re
import uuid
import zipfile
from datetime import UTC, datetime

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import func, select

from app import list_evidence
from app.auth import get_current_user
from app.db import get_session
from app.engine import (
    activate_org_product,
    deactivate_org_product,
    recompute_sprs,
    start_assessment,
)
from app.list_adoption import adopt, read_template
from app.list_library import (
    ImportRefused,
    ListVersionSnap,
    add_tag,
    apply_import,
    create_version,
)
from app.main import app
from app.models import (
    Assessment,
    AssessmentObjective,
    AuditLog,
    BaselineControl,
    Control,
    ControlState,
    ControlStateHistory,
    Evidence,
    EvidenceStateLink,
    Framework,
    ListDefinition,
    ListLinkEvent,
    Organization,
    OrgProduct,
    Product,
    ProductBaselineVersion,
)
from app.render import LIST_EXPORT_HEADER_ROW, render_list_version_bytes
from app.storage import NullStorageClient, get_storage_client
from tests.conftest import _app_session, _authed, _grant, _make_fake_user
from tests.list_template_fixtures import build_library
from tests.test_list_builder import _plan

STATE_LISTS = "/orgs/{org}/assessments/{a}/control-states/{cs}/lists"
ONE = "/orgs/{org}/list-definitions/{list}"
BUNDLE = "/orgs/{org}/assessments/{a}/bundle"

REAL_ROWS = [["2026-10-01", "A. Visitor", "J. Escort"], ["2026-10-02", "B. Visitor", "J. Escort"]]
EXAMPLE_ROWS = [["[DATE]", "[PLACEHOLDER]", "jdoe@[CUSTOMER]"], ["[DATE]", "x", "y"]]


# ---------------------------------------------------------------------------
# Pure
# ---------------------------------------------------------------------------


def _def(template=False):
    return ListDefinition(org_id=uuid.uuid4(), list_key="k", is_template=template)


class _V:
    def __init__(self, rows):
        self.rows = rows


def test_why_a_list_is_not_offered_is_said_in_words():
    assert list_evidence.not_offered_reason(_def(), _V(REAL_ROWS)) is None
    assert list_evidence.not_offered_reason(_def(), _V([])) == "it has no rows yet"
    assert list_evidence.not_offered_reason(_def(), _V(EXAMPLE_ROWS)) == (
        "still showing template examples (2 of 2 rows)"
    )
    assert list_evidence.not_offered_reason(_def(), _V([REAL_ROWS[0], EXAMPLE_ROWS[0]])) == (
        "still showing template examples (1 of 2 rows)"
    )
    assert "template" in list_evidence.not_offered_reason(_def(template=True), _V(REAL_ROWS))


def test_the_decision_is_is_populateds_not_a_second_rule(monkeypatch):
    """If is_populated were wrong, the fix belongs there. So L.4 must ask it:
    flip its answer and the offer must follow."""
    assert list_evidence.not_offered_reason(_def(), _V(EXAMPLE_ROWS)) is not None, (
        "fixture must be unpopulated or this test proves nothing"
    )
    monkeypatch.setattr(list_evidence, "is_populated", lambda rows: True)
    assert list_evidence.not_offered_reason(_def(), _V(EXAMPLE_ROWS)) is None
    monkeypatch.setattr(list_evidence, "is_populated", lambda rows: False)
    assert list_evidence.not_offered_reason(_def(), _V(REAL_ROWS)) is not None


def test_drift_is_said_the_way_l7_says_a_newer_template_version():
    assert list_evidence.drift_text(2, 2) is None
    assert list_evidence.drift_text(2, 4) == (
        "List version 4 is available. This evidence stays on version 2 until someone "
        "moves it forward."
    )


def _snap(**over):
    base = dict(
        list_id="l", list_key="visitor-log", is_template=False,
        control_keys=("PE.L2-3.10.3",), version_id="vid", version_number=3, title="Visitor Log",
        description="Everyone escorted through the CUI area.", responsible="Office Manager",
        review_cadence="Monthly", columns=("Date", "Visitor"), rows=(("2026-10-01", "A"),),
        saved_at=datetime(2026, 10, 1, 12, 0, tzinfo=UTC), saved_by="Pat",
    )
    return ListVersionSnap(**(base | over))


def test_the_export_says_what_the_list_is_for_who_owns_it_and_how_often_t():
    ws = openpyxl.load_workbook(io.BytesIO(render_list_version_bytes(_snap()))).active
    assert ws["A3"].value == "What goes here: Everyone escorted through the CUI area."
    assert ws["A4"].value == "Responsible: Office Manager   |   Review cadence: Monthly"
    assert ws.cell(row=LIST_EXPORT_HEADER_ROW, column=1).value == "Date"
    assert ws.cell(row=LIST_EXPORT_HEADER_ROW + 1, column=1).value == "2026-10-01"

    blank = openpyxl.load_workbook(io.BytesIO(render_list_version_bytes(
        _snap(description="", responsible="", review_cadence="")
    ))).active
    assert blank["A3"].value == "What goes here: not recorded"
    assert blank["A4"].value == "Responsible: not recorded   |   Review cadence: not recorded"

    typed = openpyxl.load_workbook(io.BytesIO(render_list_version_bytes(
        _snap(description="=HYPERLINK(\"http://x\")", responsible="+1", review_cadence="@now")
    ))).active
    assert [typed[c].data_type for c in ("A3", "A4")] == ["s", "s"], (
        "a typed value opens no formula: each cell starts with its label"
    )


def test_the_list_routes_that_change_a_link_are_not_gets():
    """require_write lets every GET through for c3pao_assessor; these must
    not be GETs, or the assessor could change evidence."""
    paths = app.openapi()["paths"]
    base = "/orgs/{org_id}/assessments/{assessment_id}/control-states/{cs_id}/lists"
    assert set(paths[base]) == {"get"}
    assert set(paths[base + "/{list_id}"]) == {"post", "delete"}
    assert set(paths[base + "/{list_id}/move-forward"]) == {"post"}


# ---------------------------------------------------------------------------
# Through the database and real HTTP (RLS enforced per request)
# ---------------------------------------------------------------------------


@pytest.fixture(autouse=True)
def _clear_overrides():
    yield
    app.dependency_overrides.clear()


def _world(db_session, label="w") -> dict:
    org = Organization(name=f"ListEv-{label}-{uuid.uuid4().hex[:8]}")
    fw = Framework(key=f"fw-le-{uuid.uuid4().hex[:6]}", name="LE FW", version="r2")
    db_session.add_all([org, fw])
    db_session.flush()
    controls = {}
    for i, (cid, fam, objs) in enumerate([
        ("PE.L2-3.10.3", "PE", "ab"), ("AC.L2-3.1.1", "AC", "a"), ("RA.L2-3.11.2", "RA", "a"),
    ]):
        c = Control(
            framework_id=fw.id, control_id=cid, family=fam, title=cid, requirement_text=cid,
            sprs_weight=3, sequence_order=i,
        )
        db_session.add(c)
        db_session.flush()
        controls[cid] = c
        for k in objs:
            db_session.add(
                AssessmentObjective(control_id=c.id, objective_key=k, text=f"{cid}[{k}]")
            )
    db_session.flush()
    a = start_assessment(db_session, org.id, fw.id, f"FY26 L2 {label}")
    db_session.flush()
    states = {
        f"{r.control_id}[{r.objective_key}]": r.id
        for r in db_session.execute(
            select(Control.control_id, AssessmentObjective.objective_key, ControlState.id)
            .join(AssessmentObjective, AssessmentObjective.control_id == Control.id)
            .join(ControlState, ControlState.objective_id == AssessmentObjective.id)
            .where(ControlState.assessment_id == a.id)
        )
    }
    return {"org": org.id, "fw": fw.id, "a": a.id, "cs": states, "controls": controls}


def _list(db_session, org_id, key, rows, tags=("PE.L2-3.10.3",), template=False):
    d = ListDefinition(org_id=org_id, list_key=key, is_template=template)
    db_session.add(d)
    db_session.flush()
    create_version(
        db_session, d,
        {"title": key.replace("-", " ").title(), "description": "Who came in.",
         "responsible": "Office Manager", "review_cadence": "Monthly",
         "columns": ["Date", "Visitor", "Escort"], "rows": rows},
        {"via": "api", "cause": "created_by_hand", "saved_by": "Test"},
    )
    for k in tags:
        add_tag(db_session, d, k, "api", "Test")
    db_session.flush()
    return d


def _client(db_session, org_id, role="msp_engineer") -> TestClient:
    user = _make_fake_user(
        role=role, org_id=org_id, email=f"{uuid.uuid4().hex[:8]}@x.com", display_name="Lee Linker"
    )
    _grant(db_session, user)
    app.dependency_overrides[get_session] = _app_session(db_session)
    app.dependency_overrides[get_current_user] = _authed(db_session, user)
    app.dependency_overrides[get_storage_client] = NullStorageClient
    return TestClient(app)


def _lists(client, w, cs_key):
    r = client.get(STATE_LISTS.format(org=w["org"], a=w["a"], cs=w["cs"][cs_key]))
    assert r.status_code == 200, r.text
    return {c["list_key"]: c for c in r.json()}


def _link(client, w, cs_key, list_id, action=""):
    url = f"{STATE_LISTS.format(org=w['org'], a=w['a'], cs=w['cs'][cs_key])}/{list_id}{action}"
    return client.post(url)


def _edit(client, org_id, d, rows):
    base = ONE.format(org=org_id, list=d.id)
    current = client.get(base).json()["current"]["id"]
    r = client.post(f"{base}/versions", json={"base_version_id": current, "rows": rows})
    assert r.status_code == 201, r.text
    return r.json()


@pytest.mark.integration
def test_only_a_populated_list_is_offered_and_the_rest_say_why(db_session):
    w = _world(db_session, "offer")
    _list(db_session, w["org"], "visitor-log", REAL_ROWS)
    _list(db_session, w["org"], "empty-log", [])
    _list(db_session, w["org"], "example-log", EXAMPLE_ROWS)
    _list(db_session, w["org"], "template-log", REAL_ROWS, template=True)
    _list(db_session, w["org"], "other-practice", REAL_ROWS, tags=("AC.L2-3.1.1",))
    client = _client(db_session, w["org"])

    got = _lists(client, w, "PE.L2-3.10.3[a]")
    assert set(got) == {"visitor-log", "empty-log", "example-log", "template-log"}, (
        "a list tagged with another practice is not this objective's to offer"
    )
    assert got["visitor-log"]["offered"] is True and got["visitor-log"]["not_offered"] is None
    assert got["empty-log"]["not_offered"] == "it has no rows yet"
    assert got["example-log"]["not_offered"] == "still showing template examples (2 of 2 rows)"
    assert "template" in got["template-log"]["not_offered"]

    r = _link(client, w, "PE.L2-3.10.3[a]", got["example-log"]["list_id"])
    assert r.status_code == 409 and "still showing template examples" in r.json()["detail"]


@pytest.mark.integration
def test_an_adopted_copy_is_offered_once_filled_in_with_no_other_action(db_session):
    w = _world(db_session, "adopt")
    template = _list(db_session, w["org"], "visitor-log", EXAMPLE_ROWS, template=True)
    copy = adopt(db_session, w["org"], read_template(db_session, w["org"], template.id), "Test")
    db_session.flush()
    client = _client(db_session, w["org"])

    assert _lists(client, w, "PE.L2-3.10.3[a]")[copy.list_key]["not_offered"] == (
        "still showing template examples (2 of 2 rows)"
    )
    _edit(client, w["org"], copy, REAL_ROWS)
    filled = _lists(client, w, "PE.L2-3.10.3[a]")[copy.list_key]
    assert filled["offered"] is True and filled["current_version_number"] == 2
    assert _link(client, w, "PE.L2-3.10.3[a]", copy.id).status_code == 201


@pytest.mark.integration
def test_linking_changes_no_control_state_and_no_score(db_session):
    w = _world(db_session, "noscore")
    d = _list(db_session, w["org"], "visitor-log", REAL_ROWS)
    a = db_session.get(Assessment, w["a"])
    before_score = a.sprs_score
    before = dict(db_session.execute(
        select(ControlState.id, ControlState.status).where(ControlState.assessment_id == w["a"])
    ).all())
    history = db_session.scalar(select(func.count()).select_from(ControlStateHistory))
    client = _client(db_session, w["org"])

    assert _link(client, w, "PE.L2-3.10.3[a]", d.id).status_code == 201
    assert _link(client, w, "PE.L2-3.10.3[b]", d.id).status_code == 201

    db_session.expire_all()
    assert dict(db_session.execute(
        select(ControlState.id, ControlState.status).where(ControlState.assessment_id == w["a"])
    ).all()) == before
    assert db_session.scalar(select(func.count()).select_from(ControlStateHistory)) == history
    assert db_session.get(Assessment, w["a"]).sprs_score == before_score
    assert recompute_sprs(db_session, w["a"]) == before_score
    evidence = db_session.scalars(
        select(Evidence).where(Evidence.org_id == w["org"], Evidence.artifact_type == "list")
    ).all()
    assert len(evidence) == 1, "one evidence row per list version, linked to both objectives"
    assert (evidence[0].kind, evidence[0].storage_key) == ("reference", None)


@pytest.mark.integration
def test_a_links_audit_row_names_the_objective_and_version_its_evidence_and_link_are_for(
    db_session,
):
    """The evidence row a list link creates is accounted for by a row on the
    objective (tests/audit_check.py's broad parent). This pins that the row
    describes that write: the objective the link is on and the version the
    evidence row was made from, for the first link and for a second that
    reuses the evidence."""
    w = _world(db_session, "audited")
    d = _list(db_session, w["org"], "visitor-log", REAL_ROWS)
    client = _client(db_session, w["org"])
    for key in ("PE.L2-3.10.3[a]", "PE.L2-3.10.3[b]"):
        assert _link(client, w, key, d.id).status_code == 201

    db_session.expire_all()
    ev = db_session.scalars(
        select(Evidence).where(Evidence.org_id == w["org"], Evidence.artifact_type == "list")
    ).one()
    rows = db_session.scalars(
        select(AuditLog).where(AuditLog.org_id == w["org"], AuditLog.action == "evidence.list_link")
    ).all()
    assert sorted((r.entity_type, str(r.entity_id)) for r in rows) == sorted(
        ("control_state", str(w["cs"][k])) for k in ("PE.L2-3.10.3[a]", "PE.L2-3.10.3[b]")
    )
    for r in rows:
        assert r.after_value["list_version_id"] == str(ev.source_list_version_id)
        link = db_session.scalars(
            select(EvidenceStateLink).where(EvidenceStateLink.control_state_id == r.entity_id)
        ).one()
        assert (link.evidence_id, link.is_archived) == (ev.id, False)


@pytest.mark.integration
def test_a_link_stays_on_its_version_and_the_drift_shows_in_all_three_places(db_session):
    w = _world(db_session, "drift")
    d = _list(db_session, w["org"], "visitor-log", REAL_ROWS)
    client = _client(db_session, w["org"])
    assert _link(client, w, "PE.L2-3.10.3[a]", d.id).status_code == 201
    v1 = _lists(client, w, "PE.L2-3.10.3[a]")["visitor-log"]["linked_version_id"]

    _edit(client, w["org"], d, REAL_ROWS + [["2026-10-03", "C. Visitor", "J. Escort"]])
    _edit(client, w["org"], d, REAL_ROWS)
    drift = (
        "List version 3 is available. This evidence stays on version 1 until someone "
        "moves it forward."
    )
    on_control = _lists(client, w, "PE.L2-3.10.3[a]")["visitor-log"]
    assert (on_control["linked_version_id"], on_control["linked_version_number"]) == (v1, 1)
    assert on_control["current_version_number"] == 3 and on_control["drift"] == drift

    on_list = client.get(f"{ONE.format(org=w['org'], list=d.id)}/evidence-links").json()
    assert on_list["offered"] is True
    assert [(x["control_key"], x["objective_key"], x["linked_version_number"], x["drift"])
            for x in on_list["links"]] == [("PE.L2-3.10.3", "a", 1, drift)]

    in_assessment = client.get(f"/orgs/{w['org']}/assessments/{w['a']}/list-links").json()
    assert [(x["list_key"], x["linked_version_number"], x["drift"]) for x in in_assessment] == [
        ("visitor-log", 1, drift)
    ]

    r = _link(client, w, "PE.L2-3.10.3[a]", d.id)
    assert r.status_code == 409 and "move it forward" in r.json()["detail"], (
        "a second link alongside the old one would hide the drift"
    )


@pytest.mark.integration
def test_moving_forward_writes_a_new_link_and_the_old_one_survives(db_session):
    w = _world(db_session, "forward")
    d = _list(db_session, w["org"], "visitor-log", REAL_ROWS)
    client = _client(db_session, w["org"])
    assert _link(client, w, "PE.L2-3.10.3[a]", d.id).status_code == 201
    old = db_session.scalars(
        select(EvidenceStateLink).where(
            EvidenceStateLink.control_state_id == w["cs"]["PE.L2-3.10.3[a]"]
        )
    ).one()
    old_id, old_evidence = old.id, old.evidence_id
    _edit(client, w["org"], d, REAL_ROWS[:1])

    r = _link(client, w, "PE.L2-3.10.3[a]", d.id, "/move-forward")
    assert r.status_code == 200, r.text
    moved = {c["list_key"]: c for c in r.json()}["visitor-log"]
    assert (moved["linked_version_number"], moved["drift"]) == (2, None)

    db_session.expire_all()
    links = db_session.scalars(
        select(EvidenceStateLink)
        .where(EvidenceStateLink.control_state_id == w["cs"]["PE.L2-3.10.3[a]"])
        .order_by(EvidenceStateLink.linked_at, EvidenceStateLink.id)
    ).all()
    assert len(links) == 2
    kept = db_session.get(EvidenceStateLink, old_id)
    assert kept.is_archived is True and kept.evidence_id == old_evidence, (
        "the old link is archived, never rewritten to the new version"
    )
    assert db_session.get(Evidence, old_evidence) is not None
    row = db_session.scalars(
        select(AuditLog).where(
            AuditLog.org_id == w["org"], AuditLog.action == "evidence.list_move_forward"
        )
    ).one()
    assert row.after_value["from_version_numbers"] == [1]
    assert row.after_value["version_number"] == 2

    again = _link(client, w, "PE.L2-3.10.3[a]", d.id, "/move-forward")
    assert again.status_code == 409 and "already at version 2" in again.json()["detail"]


@pytest.mark.integration
def test_a_list_cannot_be_linked_to_a_practice_it_is_not_tagged_with(db_session):
    w = _world(db_session, "untagged")
    d = _list(db_session, w["org"], "visitor-log", REAL_ROWS)
    client = _client(db_session, w["org"])
    r = _link(client, w, "AC.L2-3.1.1[a]", d.id)
    assert r.status_code == 409
    assert "not tagged with AC.L2-3.1.1" in r.json()["detail"]
    assert db_session.scalar(
        select(func.count()).select_from(Evidence).where(Evidence.org_id == w["org"])
    ) == 0


@pytest.mark.integration
def test_a_tag_a_live_link_rests_on_cannot_be_removed_by_a_person(db_session):
    """test_list_builder.py stood the link in with a monkeypatch; this is the
    real link, through tag_removal_blockers' own query."""
    w = _world(db_session, "tagperson")
    tags = ("PE.L2-3.10.3", "AC.L2-3.1.1", "RA.L2-3.11.2")
    d = _list(db_session, w["org"], "visitor-log", REAL_ROWS, tags=tags)
    client = _client(db_session, w["org"])
    base = ONE.format(org=w["org"], list=d.id)
    assert client.delete(f"{base}/tags/RA.L2-3.11.2").status_code == 200, (
        "fixture must allow removing an unlinked tag or the refusal below proves nothing"
    )
    assert _link(client, w, "PE.L2-3.10.3[b]", d.id).status_code == 201

    r = client.delete(f"{base}/tags/PE.L2-3.10.3")
    assert r.status_code == 409, r.text
    assert r.json()["detail"]["links"] == [{
        "version_number": 1, "objective_key": "PE.L2-3.10.3[b]",
        "assessment_name": "FY26 L2 tagperson",
    }], "was `blocked_by` sentences; fields since X.1, so the screen writes one sentence"
    assert "PE.L2-3.10.3" in client.get(base).json()["control_keys"]

    cs_b = w["cs"]["PE.L2-3.10.3[b]"]
    delete = client.delete(f"{STATE_LISTS.format(org=w['org'], a=w['a'], cs=cs_b)}/{d.id}")
    assert delete.status_code == 200
    assert client.delete(f"{base}/tags/PE.L2-3.10.3").status_code == 200, (
        "unlinked, the tag may go"
    )
    # This asserted the opposite until 2026-10-19 ("an unlinked list version
    # leaves no evidence row behind"): unlink deleted the link and the row.
    # A detach deletes nothing now (evidence_links.py); the tag may still go,
    # because an archived link is not one it rests on.
    assert db_session.scalar(
        select(func.count()).select_from(Evidence).where(Evidence.org_id == w["org"])
    ) == 1, "a detached list version keeps its evidence row"
    link = db_session.scalars(
        select(EvidenceStateLink).join(Evidence, Evidence.id == EvidenceStateLink.evidence_id)
        .where(Evidence.org_id == w["org"])
    ).one()
    assert (link.is_archived, link.archive_cause) == (True, "detached")
    assert link.detached_by, "says who detached it"


@pytest.mark.integration
def test_a_tag_a_live_link_rests_on_cannot_be_removed_by_an_import(db_session, tmp_path):
    w = _world(db_session, "tagimport")
    root = build_library(tmp_path)
    apply_import(db_session, w["org"], _plan(root))
    risk = db_session.scalar(
        select(ListDefinition).where(
            ListDefinition.org_id == w["org"], ListDefinition.list_key == "3-11-1a-risk-register"
        )
    )
    # As if an earlier source had named 3.11.2 and this one no longer does,
    # and the org has filled the list in (not via the API, so the import
    # treats the next run as its own and not a person's conflict).
    add_tag(db_session, risk, "RA.L2-3.11.2", "template_import", None)
    risk.is_template = False
    create_version(
        db_session, risk,
        {"title": "Risk Register", "description": "", "responsible": "", "review_cadence": "",
         "columns": ["Risk ID", "Risk Description", "Status"],
         "rows": [["R-1", "Laptop theft", "Open"]]},
        {"via": "test"},
    )
    db_session.flush()
    client = _client(db_session, w["org"])
    assert _link(client, w, "RA.L2-3.11.2[a]", risk.id).status_code == 201

    blocked = r"RA\.L2-3\.11\.2 \(version 2 linked to RA\.L2-3\.11\.2\[a\]"
    with pytest.raises(ImportRefused, match=blocked):
        apply_import(db_session, w["org"], _plan(root))
    assert "RA.L2-3.11.2" in {
        k for k in client.get(ONE.format(org=w["org"], list=risk.id)).json()["control_keys"]
    }


def _product_covering(db_session, w, control_key) -> Product:
    product = Product(
        framework_id=w["fw"], key=f"prod-{uuid.uuid4().hex[:6]}", name="Badge System",
        provider="Test Inc", category="PACS", asset_type="SPA", role="Test", is_published=True,
    )
    db_session.add(product)
    db_session.flush()
    version = ProductBaselineVersion(product_id=product.id, version_number=1)
    db_session.add(version)
    db_session.flush()
    product.current_version_id = version.id
    db_session.add(BaselineControl(
        product_id=product.id, baseline_version_id=version.id,
        control_id=w["controls"][control_key].id, objectives=["a", "b"],
        classification="shared", candidate_state="pending_evidence",
    ))
    db_session.add(OrgProduct(org_id=w["org"], product_id=product.id, status="candidate"))
    db_session.flush()
    return product


@pytest.mark.integration
def test_deactivating_a_product_on_the_same_practice_leaves_the_list_link_alone(db_session):
    w = _world(db_session, "deact")
    # Two tags: with one, L.2.2's last-tag rule refuses first and the
    # blocker below would never be asked.
    d = _list(db_session, w["org"], "visitor-log", REAL_ROWS, tags=("PE.L2-3.10.3", "AC.L2-3.1.1"))
    product = _product_covering(db_session, w, "PE.L2-3.10.3")
    activate_org_product(db_session, w["org"], product.id, w["a"])
    db_session.flush()
    cs = w["cs"]["PE.L2-3.10.3[a]"]
    tool_ev = Evidence(
        org_id=w["org"], title="Badge export", artifact_type="export", kind="reference",
        reference_location="https://badge.example/export", collected_at=datetime.now(UTC),
    )
    db_session.add(tool_ev)
    db_session.flush()
    db_session.add(EvidenceStateLink(evidence_id=tool_ev.id, control_state_id=cs))
    db_session.flush()
    client = _client(db_session, w["org"])
    assert _link(client, w, "PE.L2-3.10.3[a]", d.id).status_code == 201

    result = deactivate_org_product(db_session, w["org"], product.id, w["a"])
    db_session.flush()
    db_session.expire_all()
    archived = dict(db_session.execute(
        select(Evidence.artifact_type, EvidenceStateLink.is_archived)
        .join(Evidence, Evidence.id == EvidenceStateLink.evidence_id)
        .where(EvidenceStateLink.control_state_id == cs)
    ).all())
    assert archived == {"export": True, "list": False}, (
        "the tool's evidence is archived as before; the list a person attached is not"
    )
    assert result["evidence_links_archived"] == 1
    r = client.delete(f"{ONE.format(org=w['org'], list=d.id)}/tags/PE.L2-3.10.3")
    assert r.status_code == 409 and r.json()["detail"]["links"], (
        "the tag is still blocked: the link it rests on is still live"
    )


def _zip(content: bytes) -> zipfile.ZipFile:
    return zipfile.ZipFile(io.BytesIO(content))


def _member(zf: zipfile.ZipFile, rel: str) -> bytes:
    [name] = [n for n in zf.namelist() if n.split("/", 1)[1] == rel]
    return zf.read(name)


def _unstamped(html: bytes) -> bytes:
    return re.sub(rb"generated [^<]*", b"", html)


@pytest.mark.integration
def test_a_bundle_embeds_the_linked_version_and_says_it_is_not_current(db_session):
    w = _world(db_session, "bundle")
    d = _list(db_session, w["org"], "visitor-log", REAL_ROWS)
    _list(db_session, w["org"], "unchosen-log", REAL_ROWS)
    client = _client(db_session, w["org"])
    assert _link(client, w, "PE.L2-3.10.3[a]", d.id).status_code == 201
    unchosen = db_session.scalar(
        select(ListDefinition).where(
            ListDefinition.org_id == w["org"], ListDefinition.list_key == "unchosen-log"
        )
    )
    assert _link(client, w, "PE.L2-3.10.3[b]", unchosen.id).status_code == 201
    _edit(client, w["org"], d, REAL_ROWS[:1])

    def bundle():
        r = client.get(BUNDLE.format(org=w["org"], a=w["a"]), params={"list_id": [str(d.id)]})
        assert r.status_code == 200, r.text
        return _zip(r.content)

    zf = bundle()
    v1, v2 = _member(zf, "lists/visitor-log_v1.xlsx"), _member(zf, "lists/visitor-log_v2.xlsx")
    assert openpyxl.load_workbook(io.BytesIO(v1)).active.cell(
        row=LIST_EXPORT_HEADER_ROW + 2, column=2
    ).value == "B. Visitor", "the linked version's own rows, not the current ones"
    assert v2 == client.get(f"{ONE.format(org=w['org'], list=d.id)}/export").content

    index = _member(zf, "lists/index.html").decode()
    assert "not current</strong>; evidence for PE.L2-3.10.3[a] -- linked at version 1; the list " \
           "was at version 2" in index
    assert "Linked as evidence, not chosen for this bundle" in index
    assert "Unchosen Log, version 1" in index and "PE.L2-3.10.3[b]" in index

    manifest = _member(zf, "evidence/manifest.html").decode()
    assert 'href="../lists/visitor-log_v1.xlsx"' in manifest
    assert "list not chosen for this bundle" in manifest
    impl = _member(zf, "ssp/02_implementation.html").decode()
    assert "lists/visitor-log_v1.xlsx" in impl

    import hashlib

    log = _member(zf, "artifact_log.txt")
    [line] = [ln for ln in log.decode().splitlines() if ln.endswith("/lists/visitor-log_v1.xlsx")]
    assert line.split(" | ")[1] == hashlib.sha256(v1).hexdigest()
    assert hashlib.sha256(log).hexdigest() in _member(zf, "cover.html").decode(), (
        "the cover's second-order hash covers the log that names the linked version"
    )

    again = bundle()
    for rel in ("lists/visitor-log_v1.xlsx", "lists/visitor-log_v2.xlsx"):
        assert _member(zf, rel) == _member(again, rel)
    for rel in ("lists/index.html", "evidence/manifest.html"):
        assert _unstamped(_member(zf, rel)) == _unstamped(_member(again, rel))


@pytest.mark.integration
def test_a_list_linked_in_another_assessment_is_not_this_ones_stray_evidence(db_session):
    w = _world(db_session, "stray")
    d = _list(db_session, w["org"], "visitor-log", REAL_ROWS)
    client = _client(db_session, w["org"])
    assert _link(client, w, "PE.L2-3.10.3[a]", d.id).status_code == 201
    other = start_assessment(db_session, w["org"], w["fw"], "Second")
    db_session.flush()
    zf = _zip(client.get(BUNDLE.format(org=w["org"], a=other.id)).content)
    assert "Visitor Log" not in _member(zf, "evidence/manifest.html").decode()


@pytest.mark.integration
def test_an_assessor_sees_links_and_can_make_move_or_break_none(db_session):
    w = _world(db_session, "c3pao")
    d = _list(db_session, w["org"], "visitor-log", REAL_ROWS)
    engineer = _client(db_session, w["org"])
    assert _link(engineer, w, "PE.L2-3.10.3[a]", d.id).status_code == 201
    _edit(engineer, w["org"], d, REAL_ROWS[:1])

    assessor = _client(db_session, w["org"], role="c3pao_assessor")
    assert _lists(assessor, w, "PE.L2-3.10.3[a]")["visitor-log"]["linked_version_number"] == 1
    assert assessor.get(f"/orgs/{w['org']}/assessments/{w['a']}/list-links").status_code == 200
    assert assessor.get(f"{ONE.format(org=w['org'], list=d.id)}/evidence-links").status_code == 200
    assert _link(assessor, w, "PE.L2-3.10.3[b]", d.id).status_code == 403
    assert _link(assessor, w, "PE.L2-3.10.3[a]", d.id, "/move-forward").status_code == 403
    url = f"{STATE_LISTS.format(org=w['org'], a=w['a'], cs=w['cs']['PE.L2-3.10.3[a]'])}/{d.id}"
    assert assessor.delete(url).status_code == 403
    assert assessor.get(f"{ONE.format(org=w['org'], list=d.id)}/export").status_code == 200, (
        "the export is a GET and must stay one (the lists slice's section 0)"
    )
    assert db_session.scalar(
        select(func.count()).select_from(EvidenceStateLink)
        .where(EvidenceStateLink.control_state_id.in_(list(w["cs"].values())))
    ) == 1, "the assessor's attempts wrote nothing"


@pytest.mark.integration
def test_no_org_reaches_another_orgs_list_links(db_session):
    a = _world(db_session, "rls-a")
    b = _world(db_session, "rls-b")
    d = _list(db_session, a["org"], "visitor-log", REAL_ROWS)
    owner = _client(db_session, a["org"])
    assert _link(owner, a, "PE.L2-3.10.3[a]", d.id).status_code == 201

    intruder = _client(db_session, b["org"])
    assert intruder.get(
        STATE_LISTS.format(org=a["org"], a=a["a"], cs=a["cs"]["PE.L2-3.10.3[a]"])
    ).status_code == 403, "no membership in org A"
    # Org B's own path, org A's ids: RLS hides them, so not found.
    assert intruder.get(
        STATE_LISTS.format(org=b["org"], a=a["a"], cs=a["cs"]["PE.L2-3.10.3[a]"])
    ).status_code == 404
    assert _link(intruder, b, "PE.L2-3.10.3[a]", d.id).status_code == 404
    assert intruder.get(f"{ONE.format(org=b['org'], list=d.id)}/evidence-links").status_code == 404
    assert intruder.get(f"/orgs/{b['org']}/assessments/{b['a']}/list-links").json() == []
    assert _lists(intruder, b, "PE.L2-3.10.3[a]") == {}


# ---------------------------------------------------------------------------
# L.10: what the screens say is computed, and the list's history has the links
# ---------------------------------------------------------------------------


@pytest.mark.integration
def test_the_removal_check_says_what_the_delete_would_do(db_session):
    """The tag dialog read the rule aloud ("it is refused while evidence
    rests on this list") beside an enabled button, whatever the links were.
    It now asks this, which runs the DELETE's own checks.

    Was `refused`/`blocked_by` sentences, built for a log: on live the
    dialog wrapped them in its own prefix and named the list by key, one
    fact three times (L.10's live finding). Since X.1 the check answers
    in fields, and the DELETE's 409 carries the same fields."""
    w = _world(db_session, "check")
    tags = ("PE.L2-3.10.3", "AC.L2-3.1.1")
    d = _list(db_session, w["org"], "visitor-log", REAL_ROWS, tags=tags)
    client = _client(db_session, w["org"])
    base = ONE.format(org=w["org"], list=d.id)

    assert client.get(f"{base}/tags/PE.L2-3.10.3/removal").json() == {
        "removable": True, "reason": None, "links": [],
    }
    assert _link(client, w, "PE.L2-3.10.3[a]", d.id).status_code == 201
    check = client.get(f"{base}/tags/PE.L2-3.10.3/removal").json()
    refused = client.delete(f"{base}/tags/PE.L2-3.10.3")
    assert refused.status_code == 409
    links = [{
        "version_number": 1, "objective_key": "PE.L2-3.10.3[a]",
        "assessment_name": "FY26 L2 check",
    }]
    assert check == {"removable": False, "reason": "evidence", "links": links}
    detail = refused.json()["detail"]
    assert (detail["reason"], detail["links"]) == ("evidence", links), (
        "the check and the DELETE must say the same thing"
    )

    assert client.delete(f"{base}/tags/AC.L2-3.1.1").status_code == 200
    last = client.get(f"{base}/tags/PE.L2-3.10.3/removal").json()
    assert last == {"removable": False, "reason": "last_tag", "links": []}
    assert client.get(f"{base}/tags/RA.L2-3.11.2/removal").status_code == 404


@pytest.mark.integration
def test_the_lists_history_answers_was_this_ever_attached(db_session):
    """On live (2026-10-18) the question took the request log to answer:
    the list's history had versions and tags and no links. Link, move
    forward and unlink are now there, and stay after the link is gone."""
    w = _world(db_session, "history")
    tags = ("PE.L2-3.10.3", "AC.L2-3.1.1")
    d = _list(db_session, w["org"], "visitor-log", REAL_ROWS, tags=tags)
    client = _client(db_session, w["org"])
    base = ONE.format(org=w["org"], list=d.id)
    assert client.get(f"{base}/link-events").json() == [], (
        "fixture must start with no events or this test proves nothing"
    )

    assert _link(client, w, "PE.L2-3.10.3[a]", d.id).status_code == 201
    _edit(client, w["org"], d, REAL_ROWS[:1])
    assert _link(client, w, "PE.L2-3.10.3[a]", d.id, "/move-forward").status_code == 200
    cs_a = w["cs"]["PE.L2-3.10.3[a]"]
    assert client.delete(
        f"{STATE_LISTS.format(org=w['org'], a=w['a'], cs=cs_a)}/{d.id}"
    ).status_code == 200

    events = client.get(f"{base}/link-events").json()
    assert [
        (e["change"], e["version_number"], e["from_version_numbers"], e["control_key"],
         e["objective_key"], e["assessment_name"], e["saved_by"])
        for e in events
    ] == [
        ("linked", 1, None, "PE.L2-3.10.3", "a", "FY26 L2 history", "Lee Linker"),
        ("moved_forward", 2, [1], "PE.L2-3.10.3", "a", "FY26 L2 history", "Lee Linker"),
        ("unlinked", 2, None, "PE.L2-3.10.3", "a", "FY26 L2 history", "Lee Linker"),
    ]
    assert client.delete(f"{base}/tags/PE.L2-3.10.3").status_code == 200, (
        "an event is a record, not a link: an unlinked list's history blocks no tag"
    )

    consultant = _client(db_session, w["org"], role="consultant_admin")
    assert len(consultant.get(f"{base}/link-events").json()) == 3, (
        "the reason this is a table: a role that cannot read the audit log reads this"
    )

    event = db_session.scalars(select(ListLinkEvent).where(ListLinkEvent.list_id == d.id)).first()
    with pytest.raises(Exception, match="append-only"):
        with db_session.begin_nested():
            event.saved_by = "someone else"
            db_session.flush()
