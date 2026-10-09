"""Lists follow-up (2026-10-07), Jarrod's three answers:

1. Out-of-boundary entities are excluded from the lists -- they were never
   authorized -- and the exclusion is stated on screen and in the sheet.
   Decommissioned entities stay, with their date.
2. Workbook imports honour the same boundary/status invariant Liongard
   syncs do: no import changes an existing row's in_boundary, and none
   brings a decommissioned row back.
3. "CUI", "cui asset", "Specialized", "out-of-scope"... resolve to their
   ScopeCategory. Plus the general guard: export a list, re-import the
   export, nothing changes.
"""

from __future__ import annotations

import logging
import uuid
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

from app.auth import get_current_user
from app.catalog import ALL_VIEWS, AUTHORIZED_DEVICES
from app.db import get_session
from app.domain import CanonicalEntity, EntityStatus, EntityType, ScopeCategory, Source
from app.importers.workbook import parse_workbook, resolve_category
from app.list_projection import excluded_count, project
from app.main import app
from app.models import Organization, ScopeEntity
from app.render import render_view_bytes
from tests.conftest import _app_session, _authed, _grant, _make_fake_user

_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _device(key: str, *, in_boundary=True, status=EntityStatus.ACTIVE, **attrs) -> CanonicalEntity:
    return CanonicalEntity(
        entity_type=EntityType.DEVICE,
        natural_key=key,
        attributes={"Name": key, "Serial # or Asset Tag": key, **attrs},
        in_boundary=in_boundary,
        status=status,
    )


# ---------------------------------------------------------------------------
# 5.1 -- exclusion, stated
# ---------------------------------------------------------------------------


def test_out_of_boundary_is_excluded_and_decommissioned_kept_with_its_date():
    entities = [
        _device("IN-1"),
        _device("REJECTED-1", in_boundary=False),
        _device(
            "OLD-1",
            status=EntityStatus.DECOMMISSIONED,
            decommissioned_date="2025-06-30",
        ),
    ]
    rows = project(AUTHORIZED_DEVICES, entities)
    assert [r.natural_key for r in rows] == ["IN-1", "OLD-1"]
    cols = [k for k, _ in AUTHORIZED_DEVICES.columns]
    assert rows[1].cells[cols.index("Decommissioned Date")].value == "2025-06-30"
    assert excluded_count(AUTHORIZED_DEVICES, entities) == 1


def test_the_sheet_states_the_exclusion():
    ws = openpyxl.load_workbook(
        BytesIO(
            render_view_bytes(
                AUTHORIZED_DEVICES,
                [
                    _device("IN-1"),
                    _device("R-1", in_boundary=False),
                    _device("R-2", in_boundary=False),
                ],
            )
        )
    ).active
    assert "2 entities excluded as out of the CUI boundary" in ws["A3"].value
    assert "(1 record(s))" in ws["A3"].value
    values = {c.value for row in ws.iter_rows(min_row=6) for c in row}
    assert "R-1" not in values and "R-2" not in values


def test_no_exclusion_note_when_nothing_is_excluded():
    ws = openpyxl.load_workbook(
        BytesIO(render_view_bytes(AUTHORIZED_DEVICES, [_device("IN-1")]))
    ).active
    assert "excluded" not in ws["A3"].value


# ---------------------------------------------------------------------------
# 5.3 -- category aliases
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("CUI", ScopeCategory.CUI_ASSET),
        ("cui asset", ScopeCategory.CUI_ASSET),
        ("  CUI   Asset ", ScopeCategory.CUI_ASSET),
        ("CRMA", ScopeCategory.CRMA),
        ("crma asset", ScopeCategory.CRMA),
        ("SPA", ScopeCategory.SPA),
        ("Spa Asset", ScopeCategory.SPA),
        ("Specialized", ScopeCategory.SPECIALIZED),
        ("specialized asset", ScopeCategory.SPECIALIZED),
        ("Out of Scope", ScopeCategory.OUT_OF_SCOPE),
        ("out of scope asset", ScopeCategory.OUT_OF_SCOPE),
        ("out-of-scope", ScopeCategory.OUT_OF_SCOPE),
        ("ESP", ScopeCategory.ESP),
        ("csp", ScopeCategory.CSP),
        ("Unclassified", ScopeCategory.UNCLASSIFIED),
        ("Security Protection Asset", ScopeCategory.SPA),
        ("Contractor Risk Managed Asset", ScopeCategory.CRMA),
    ],
)
def test_category_aliases_resolve(raw, expected):
    assert resolve_category(raw) == expected


@pytest.mark.parametrize("raw", ["Laptop", "", None, "CUI-ish", 5])
def test_unknown_categories_do_not_resolve(raw):
    assert resolve_category(raw) is None


def test_alias_matches_are_reported_per_row_and_logged_at_warning(tmp_path, caplog, monkeypatch):
    """Each alias is reported against its row (the dry-run shows it to the
    person confirming the import) and logged at WARNING. The first version
    logged at INFO, which nothing in the app configures -- so on live it was
    silently dropped, and only a live check found that out (2026-10-09)."""
    # conftest's Alembic run calls fileConfig(), whose default
    # disable_existing_loggers=True switches off every logger already
    # imported -- this one included -- for the rest of the session. A
    # suite-only artifact: production runs migrations in its own process.
    monkeypatch.setattr(logging.getLogger("app.importers.workbook"), "disabled", False)
    notes: dict = {}
    with caplog.at_level(logging.WARNING, logger="app.importers.workbook"):
        parse_workbook(_jarrod_style_workbook(tmp_path / "w.xlsx"), notes=notes)
    assert notes[("device", "sn-dt26")] == ["Asset Type 'CUI' read as 'CUI Asset'"]
    assert notes[("device", "sn-fw01")] == ["Asset Type 'spa' read as 'SPA'"]
    assert notes[("device", "sn-prn01")] == ["Asset Type 'crma asset' read as 'CRMA'"]
    assert notes[("external_service", "liongard")] == ["Asset Type 'Spa' read as 'SPA'"]
    assert len(notes) == 4, "a canonical spelling must not be reported"
    assert [r.levelno for r in caplog.records] == [logging.WARNING] * 4
    assert "'CUI' read as 'CUI Asset'" in caplog.records[0].getMessage()


def _jarrod_style_workbook(path: Path) -> Path:
    """A workbook spelled the way Jarrod's real one is: CUI/CRMA/SPA in
    lower and mixed case, and a [PLACEHOLDER - reason] cell."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for view in ALL_VIEWS:
        ws = wb.create_sheet(view.sheet_title)
        ws["A1"] = "Authorized Entities"
        ws.append([])
        ws.append([display for _, display in view.columns])
    dev = wb["3.1.1c Authorized Devices"]
    cols = [d for _, d in AUTHORIZED_DEVICES.columns]

    def row(**v):
        dev.append([v.get(c) for c in cols])

    row(
        **{
            "Name": "WL-DT26",
            "Make": "ASUS",
            "Serial # or Asset Tag": "SN-DT26",
            "OS": "Windows 11",
            "BIOS FW Ver": "[PLACEHOLDER - BIOS FW version not collected]",
            "Location": "Office",
            "Asset Type": "CUI",
            "In Service Date": "2025-01-10",
        }
    )
    row(
        **{
            "Name": "FW-01",
            "Make": "Fortinet",
            "Serial # or Asset Tag": "SN-FW01",
            "Asset Type": "spa",
        }
    )
    row(
        **{
            "Name": "PRN-01",
            "Make": "HP",
            "Serial # or Asset Tag": "SN-PRN01",
            "Asset Type": "crma asset",
        }
    )
    wb["External Services"].append(["Liongard", "ESP", "Spa"])
    wb["3.1.1b Auth Processes"].append(
        ["Liongard inspector agent", "SRV-01", "svc-liongard", "Inventory"]
    )
    wb["3.1.1a Authorized Users"].append(
        ["Ada", "Lovelace", "Yes", "Yes", "Security Officer", "2025-01-06", None]
    )
    wb.save(path)
    return path


def test_jarrod_style_categories_parse(tmp_path):
    entities = {
        e.natural_key: e for e in parse_workbook(_jarrod_style_workbook(tmp_path / "w.xlsx"))
    }
    assert entities["SN-DT26"].scope_category == ScopeCategory.CUI_ASSET
    assert entities["SN-FW01"].scope_category == ScopeCategory.SPA
    assert entities["SN-PRN01"].scope_category == ScopeCategory.CRMA
    assert entities["Liongard"].scope_category == ScopeCategory.SPA
    assert entities["SN-DT26"].attributes["Asset Type"] == "CUI Asset"
    assert (
        entities["SN-DT26"].attributes["BIOS FW Ver"]
        == "[PLACEHOLDER - BIOS FW version not collected]"
    )


# ---------------------------------------------------------------------------
# Real HTTP, real RLS
# ---------------------------------------------------------------------------


@pytest.fixture(autouse=True)
def _clear_overrides():
    yield
    app.dependency_overrides.clear()


def _client(db_session) -> tuple[TestClient, uuid.UUID]:
    org = Organization(name=f"ListsFollowupOrg-{uuid.uuid4().hex[:8]}")
    db_session.add(org)
    db_session.flush()
    user = _make_fake_user(
        role="msp_engineer", org_id=org.id, email=f"{uuid.uuid4().hex[:8]}@example.com"
    )
    _grant(db_session, user)
    app.dependency_overrides[get_session] = _app_session(db_session)
    app.dependency_overrides[get_current_user] = _authed(db_session, user)
    return TestClient(app), org.id


def _dry_run(client, org_id, name: str, data: bytes) -> list[dict]:
    r = client.post(f"/orgs/{org_id}/imports/workbook/dry-run", files={"file": (name, data, _XLSX)})
    assert r.status_code == 200, r.text
    return r.json()["changes"]


def _import(client, org_id, name: str, data: bytes) -> None:
    changes = [
        c for c in _dry_run(client, org_id, name, data) if c["change_type"] in ("new", "changed")
    ]
    r = client.post(f"/orgs/{org_id}/imports/workbook/apply", json={"changes": changes})
    assert r.status_code == 200, r.text


def _row(db_session, org_id, key) -> ScopeEntity:
    db_session.expire_all()
    return db_session.scalars(
        select(ScopeEntity).where(ScopeEntity.org_id == org_id, ScopeEntity.natural_key == key)
    ).one()


@pytest.mark.integration
def test_export_then_reimport_changes_nothing(db_session, tmp_path):
    """The general guard: whatever a list renders must read back as the
    same entity. Catches a spelling mismatch (CUI vs CUI Asset) or any other
    projection/import disagreement without enumerating them."""
    client, org_id = _client(db_session)
    _import(client, org_id, "jarrod.xlsx", _jarrod_style_workbook(tmp_path / "j.xlsx").read_bytes())
    sample = Path(__file__).resolve().parents[2] / "samples" / "authorized-entities.example.xlsx"
    _import(client, org_id, sample.name, sample.read_bytes())

    for view in ALL_VIEWS:
        exported = client.get(f"/orgs/{org_id}/exports/{view.id}")
        assert exported.status_code == 200
        changes = _dry_run(client, org_id, f"{view.id}.xlsx", exported.content)
        edits = [
            (c["change_type"], c["natural_key"], c["field_diffs"])
            for c in changes
            if c["change_type"] in ("new", "changed") and c["entity_type"] == view.entity_type.value
        ]
        assert edits == [], f"{view.id}: re-importing its own export proposed {edits}"


@pytest.mark.integration
def test_dry_run_shows_the_alias_on_the_row(db_session, tmp_path):
    client, org_id = _client(db_session)
    changes = _dry_run(
        client, org_id, "j.xlsx", _jarrod_style_workbook(tmp_path / "j.xlsx").read_bytes()
    )
    row = next(c for c in changes if c["natural_key"] == "SN-DT26")
    assert row["incoming"]["scope_category"] == "CUI Asset"
    assert "Asset Type 'CUI' read as 'CUI Asset'" in row["warnings"]


@pytest.mark.integration
def test_workbook_reimport_keeps_a_rejected_device_out_of_scope(db_session, tmp_path):
    client, org_id = _client(db_session)
    data = _jarrod_style_workbook(tmp_path / "j.xlsx").read_bytes()
    _import(client, org_id, "j.xlsx", data)
    row = _row(db_session, org_id, "SN-FW01")
    assert (
        client.patch(f"/orgs/{org_id}/scope/{row.id}", json={"in_boundary": False}).status_code
        == 200
    )

    wb = openpyxl.load_workbook(BytesIO(data))
    ws = wb["3.1.1c Authorized Devices"]
    for r in ws.iter_rows():
        if r[0].value == "FW-01":
            r[2].value = "Fortinet Inc."  # a real edit, so the row is applied as CHANGED
    buf = BytesIO()
    wb.save(buf)
    _import(client, org_id, "j2.xlsx", buf.getvalue())

    row = _row(db_session, org_id, "SN-FW01")
    assert row.attributes["Make"] == "Fortinet Inc.", "the import itself must still land"
    assert row.in_boundary is False


@pytest.mark.integration
def test_workbook_reimport_does_not_bring_back_a_decommissioned_device(db_session, tmp_path):
    client, org_id = _client(db_session)
    data = _jarrod_style_workbook(tmp_path / "j.xlsx").read_bytes()
    _import(client, org_id, "j.xlsx", data)
    row = _row(db_session, org_id, "SN-PRN01")
    client.patch(f"/orgs/{org_id}/scope/{row.id}", json={"status": "decommissioned"})

    wb = openpyxl.load_workbook(BytesIO(data))
    for r in wb["3.1.1c Authorized Devices"].iter_rows():
        if r[0].value == "PRN-01":
            r[2].value = "HP Inc."
    buf = BytesIO()
    wb.save(buf)
    _import(client, org_id, "j2.xlsx", buf.getvalue())
    assert _row(db_session, org_id, "SN-PRN01").status == "decommissioned"


@pytest.mark.integration
def test_a_workbook_can_still_decommission_via_its_date_column(db_session, tmp_path):
    client, org_id = _client(db_session)
    data = _jarrod_style_workbook(tmp_path / "j.xlsx").read_bytes()
    _import(client, org_id, "j.xlsx", data)

    wb = openpyxl.load_workbook(BytesIO(data))
    cols = [d for _, d in AUTHORIZED_DEVICES.columns]
    for r in wb["3.1.1c Authorized Devices"].iter_rows():
        if r[0].value == "PRN-01":
            r[cols.index("Decommissioned Date")].value = "2026-09-30"
    buf = BytesIO()
    wb.save(buf)
    _import(client, org_id, "j2.xlsx", buf.getvalue())
    assert _row(db_session, org_id, "SN-PRN01").status == "decommissioned"


@pytest.mark.integration
def test_lists_endpoints_report_the_exclusion(db_session, tmp_path):
    client, org_id = _client(db_session)
    _import(client, org_id, "j.xlsx", _jarrod_style_workbook(tmp_path / "j.xlsx").read_bytes())
    row = _row(db_session, org_id, "SN-FW01")
    client.patch(f"/orgs/{org_id}/scope/{row.id}", json={"in_boundary": False})

    body = client.get(f"/orgs/{org_id}/lists/{AUTHORIZED_DEVICES.id}").json()
    assert sorted(r["natural_key"] for r in body["rows"]) == ["SN-DT26", "SN-PRN01"]
    assert body["excluded_out_of_boundary"] == 1
    assert body["excluded_note"].startswith("1 entity excluded as out of the CUI boundary")

    summary = {v["id"]: v for v in client.get(f"/orgs/{org_id}/lists").json()}
    assert summary[AUTHORIZED_DEVICES.id]["row_count"] == 2
    assert summary[AUTHORIZED_DEVICES.id]["excluded_out_of_boundary"] == 1


@pytest.mark.integration
def test_liongard_and_workbook_writes_both_leave_in_boundary_alone(db_session):
    """Direct at the choke point, both sources side by side."""
    from app import repo

    org = Organization(name=f"ListsFollowupRepo-{uuid.uuid4().hex[:8]}")
    db_session.add(org)
    db_session.flush()
    for source in (Source.WORKBOOK, Source.LIONGARD):
        key = f"K-{source.value}"
        repo.upsert(db_session, org.id, _device(key), operator_edit=True)
        repo.upsert(db_session, org.id, _device(key, in_boundary=False), operator_edit=True)
        incoming = _device(key)
        incoming.source = source
        repo.upsert(db_session, org.id, incoming)
        db_session.flush()
        assert _row(db_session, org.id, key).in_boundary is False, source


def test_reconcile_surfaces_an_import_asserted_decommission():
    """A workbook row whose only change is gaining a Decommissioned Date
    used to reconcile as UNCHANGED (status was never compared), so it was
    never applied. A sync's default status=active must still not register."""
    from app.domain import ChangeType
    from app.reconcile import reconcile

    current = _device("OLD-1")
    gone = _device("OLD-1", status=EntityStatus.DECOMMISSIONED)
    [change] = reconcile([current], [gone]).changes
    assert change.change_type == ChangeType.CHANGED
    assert change.field_diffs["status"] == ("active", "decommissioned")

    already = _device("OLD-1", status=EntityStatus.DECOMMISSIONED)
    [change] = reconcile([already], [_device("OLD-1")]).changes
    assert change.change_type == ChangeType.UNCHANGED
