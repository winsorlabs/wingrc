"""Synthetic template workbooks in the exact shape of Jarrod's library
(docs/PLAN-list-library.md §3), surveyed 2026-10-10: the title and practice
rows are merged cells whose text sits in column C, the header block runs
rows 1-9, columns are row 10, rows start at 11. The real library is MSP
content and is not committed; these reproduce its conventions, including
the irregular shapes the importer must report rather than resolve."""

from __future__ import annotations

from pathlib import Path

import openpyxl

DEVICE_COLUMNS = [
    "Name",
    "Owner / Primary User",
    "Make",
    "Model",
    "Serial # or Asset Tag",
    "Mac Address",
    "OS",
    "BIOS FW Ver",
    "Location",
    "Asset Type",
    "In Service Date",
    "Decommissioned Date",
    "FenixPyre Installed",
    "DUO Installed",
    "Senteon Installed",
    "RoboShadow Installed",
    "Heimdal Installed",
]
HARDWARE_COLUMNS = [
    "Asset Name",
    "Type",
    "Owner / User",
    "Make / Model",
    "Serial / Asset Tag",
    "OS / Firmware",
    "Baseline Ref",
    "Location",
]


def template_sheet(
    ws,
    *,
    title: str,
    practices: str,
    columns: list[str],
    rows: list[list[str]] = (),
    section: str = "Section",
    description: str = "What goes here: things.",
    responsible: str = "Responsible: MSP Engineer    |    Review cadence: Quarterly",
) -> None:
    width = max(len(columns), 3)
    last = openpyxl.utils.get_column_letter(width)
    ws["C1"], ws["C2"] = title, f"CMMC Practice / NIST 800-171: {practices}"
    ws.merge_cells(f"C1:{last}1")
    ws.merge_cells(f"C2:{last}2")
    ws["A3"], ws["A4"], ws["A5"] = "HOW TO USE THIS TEMPLATE", description, responsible
    ws["A6"] = "Replace every [PLACEHOLDER] with customer-specific data before use."
    ws["A8"], ws["C8"] = "Updated Date:", "Updated By:"
    ws["A9"] = section
    for i, c in enumerate(columns, start=1):
        ws.cell(row=10, column=i, value=c)
    for r, row in enumerate(rows, start=11):
        for i, v in enumerate(row, start=1):
            ws.cell(row=r, column=i, value=v)


def _book(path: Path, sheets: list[tuple[str, dict]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, kw in sheets:
        template_sheet(wb.create_sheet(name), **kw)
    wb.save(path)


def build_library(root: Path, *, risk_status: str = "Open") -> Path:
    """A miniature library under `root`: New Lists, Changelog/Lists, Archive/Lists."""
    nl = root / "New Lists"
    _book(
        nl / "RA/3.11.1/Risk Register.xlsx",
        [
            (
                "3.11.1a Risk Register",
                dict(
                    title="Risk Register",
                    practices="RA.L2-3.11.1",
                    section="Organizational Risk Register",
                    description="What goes here: Identified risks to CUI.",
                    responsible="Responsible: Customer Risk Owner    |    Review cadence: Annually",
                    columns=["Risk ID", "Risk Description", "Status"],
                    rows=[
                        ["[PLACEHOLDER]", "[PLACEHOLDER - no risks assessed yet]", risk_status],
                        ["[PLACEHOLDER]", "[PLACEHOLDER - no risks assessed yet]", risk_status],
                    ],
                ),
            )
        ],
    )
    _book(
        nl / "CA/3.12.2/POA&M.xlsx",
        [
            (
                "3.12.2a POAM",
                dict(
                    title="POA&M",
                    practices="CA.L2-3.12.1 / 3.12.2 / 3.12.3",
                    columns=["Weakness", "Milestone"],
                    rows=[["W1", "M1"]],
                ),
            )
        ],
    )
    _book(
        nl / "AC/3.1.1/Authorized-Entities.xlsx",
        [
            (
                "3.1.1a Authorized Users",
                dict(
                    title="Authorized Entities",
                    practices="AC.L2-3.1.1",
                    columns=["First Name", "Last Name"],
                    rows=[["Ada", "Lovelace"]],
                    section="Authorized Users - Active Directory",
                ),
            ),
            (
                "3.1.1c Authorized Devices",
                dict(
                    title="Authorized Entities",
                    practices="AC.L2-3.1.1",
                    columns=DEVICE_COLUMNS,
                    rows=[["[PLACEHOLDER]"] * len(DEVICE_COLUMNS)],
                ),
            ),
        ],
    )
    # A second section under the same tab header: "the tool is a column value".
    wb = openpyxl.load_workbook(nl / "AC/3.1.1/Authorized-Entities.xlsx")
    ws = wb["3.1.1a Authorized Users"]
    ws["A13"], ws["A14"], ws["B14"] = "Authorized Users - Datto RMM", "First Name", "Last Name"
    ws["A15"], ws["B15"] = "Ada", "Lovelace"
    wb.save(nl / "AC/3.1.1/Authorized-Entities.xlsx")
    _book(
        nl / "CM/3.4.1/Hardware and Software Asset Inventory.xlsx",
        [
            (
                "3.4.1a Hardware",
                dict(
                    title="Hardware & Software Asset Inventory",
                    practices="CM.L2-3.4.1 / 3.4.2",
                    columns=HARDWARE_COLUMNS,
                    rows=[["[PLACEHOLDER]", "Workstation"] + [""] * 6],
                ),
            ),
            (
                "3.4.1b Software",
                dict(
                    title="Hardware & Software Asset Inventory",
                    practices="CM.L2-3.4.1 / 3.4.2",
                    columns=["Software / Application", "Version"],
                    rows=[["[PLACEHOLDER]", "1"]],
                ),
            ),
        ],
    )
    _book(
        nl / "AC/3.1.5/Priviledged accounts.xlsx",
        [
            (
                "3.1.5c",
                dict(
                    title="Privileged accounts",
                    practices="AC.L2-3.1.5",
                    columns=["Function Type", "Description", "Related Systems"] * 2,
                ),
            )
        ],
    )
    # Non-conforming: no header block at all.
    path = nl / "MA/CMMC_372_Controls.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.active["A1"] = "Control"
    wb.save(path)
    path = nl / "AC/3.1.20/CMMC_3.1.20_External_Connections.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.active["A1"] = "External connections"
    wb.save(path)

    cl = root / "Changelog" / "Lists"
    cl.mkdir(parents=True, exist_ok=True)
    (cl / "New_Lists_changelog.md").write_text(
        "# New Lists\n\n- Risk Register: replaced generic example rows.\n", encoding="utf-8"
    )
    ar = root / "Archive" / "Lists" / "AC" / "3.1.1"
    ar.mkdir(parents=True, exist_ok=True)
    openpyxl.Workbook().save(ar / "Authorized-Entities.xlsx")
    return root
