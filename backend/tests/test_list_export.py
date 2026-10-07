"""List export (`GET /orgs/{org_id}/exports/{view_id}`): tenant isolation
and read-only access.

Two defects found by reading `routers/scope.py` while grounding the Lists
slice, not by a test or a user report:

- The export rendered to `tempdir/{view_id}.xlsx` -- one path per view,
  shared by every org -- and returned a `FileResponse` that read that path
  only when the response was sent. A second org's export landing between
  the first one's render and its send overwrote the file the first was
  about to serve. Fixed by rendering in memory; there is no file to share.
- The export was a POST, so `require_write()` 403'd `c3pao_assessor`.

The interleaving test runs without a database: it calls the handler
directly, the way the race interleaved it (A's handler returns, B's runs,
then A's body is read), with `repo.list_entities` stubbed per org.
"""
from __future__ import annotations

import uuid
from io import BytesIO

import openpyxl
import pytest
from fastapi.testclient import TestClient
from starlette.responses import FileResponse

from app.auth import get_current_user
from app.catalog import ALL_VIEWS, AUTHORIZED_DEVICES, VIEWS_BY_ID
from app.db import get_session
from app.domain import CanonicalEntity, EntityType
from app.main import app
from app.models import Organization, ScopeEntity
from app.routers import scope as scope_router
from tests.conftest import _app_session, _authed, _grant, _make_fake_user


def _device(name: str) -> CanonicalEntity:
    return CanonicalEntity(
        entity_type=EntityType.DEVICE,
        natural_key=name,
        attributes={"Name": name, "Serial # or Asset Tag": f"SN-{name}"},
    )


def _body(resp) -> bytes:
    # The pre-fix handler returned a FileResponse whose bytes were read
    # from disk only at send time; reading the path here is what sending
    # it would have done.
    if isinstance(resp, FileResponse):
        with open(resp.path, "rb") as fh:
            return fh.read()
    return resp.body


def _cell_values(xlsx: bytes) -> set[str]:
    ws = openpyxl.load_workbook(BytesIO(xlsx)).active
    return {str(c.value) for row in ws.iter_rows() for c in row if c.value is not None}


def test_interleaved_exports_of_one_view_each_serve_their_own_org(monkeypatch):
    org_a, org_b = uuid.uuid4(), uuid.uuid4()
    rows = {org_a: [_device("ALPHA-LT-01")], org_b: [_device("BRAVO-LT-01")]}
    monkeypatch.setattr(
        scope_router.repo, "list_entities", lambda _s, org_id, _t=None: rows[org_id]
    )
    view_id = AUTHORIZED_DEVICES.id

    resp_a = scope_router.export_view(org_a, view_id, session=None)
    resp_b = scope_router.export_view(org_b, view_id, session=None)

    cells_a, cells_b = _cell_values(_body(resp_a)), _cell_values(_body(resp_b))
    assert "ALPHA-LT-01" in cells_a and "BRAVO-LT-01" not in cells_a
    assert "BRAVO-LT-01" in cells_b and "ALPHA-LT-01" not in cells_b


def test_export_writes_no_file(monkeypatch, tmp_path):
    monkeypatch.setattr(scope_router.repo, "list_entities", lambda *_a, **_k: [])
    monkeypatch.setattr("tempfile.tempdir", str(tmp_path))
    scope_router.export_view(uuid.uuid4(), AUTHORIZED_DEVICES.id, session=None)
    assert list(tmp_path.iterdir()) == []


def test_export_route_is_get():
    ops = app.openapi()["paths"]["/orgs/{org_id}/exports/{view_id}"]
    assert set(ops) == {"get"}


# ---------------------------------------------------------------------------
# Real HTTP, real RLS (wingrc_app via _app_session)
# ---------------------------------------------------------------------------


@pytest.fixture(autouse=True)
def _clear_overrides():
    yield
    app.dependency_overrides.clear()


def _seed_org_with_device(db_session, device_name: str) -> uuid.UUID:
    org = Organization(name=f"ListExportOrg-{uuid.uuid4().hex[:8]}")
    db_session.add(org)
    db_session.flush()
    db_session.add(
        ScopeEntity(
            org_id=org.id,
            entity_type=EntityType.DEVICE.value,
            natural_key=device_name,
            status="active",
            in_boundary=True,
            source="manual",
            attributes={"Name": device_name},
        )
    )
    db_session.flush()
    return org.id


def _client(db_session, user) -> TestClient:
    app.dependency_overrides[get_session] = _app_session(db_session)
    app.dependency_overrides[get_current_user] = _authed(db_session, user)
    return TestClient(app)


@pytest.mark.integration
def test_two_orgs_export_same_view_over_http_each_get_only_their_own(db_session):
    org_a = _seed_org_with_device(db_session, "ALPHA-WS-01")
    org_b = _seed_org_with_device(db_session, "BRAVO-WS-01")
    user_a = _make_fake_user(role="msp_engineer", org_id=org_a, email="a@example.com")
    user_b = _make_fake_user(role="msp_engineer", org_id=org_b, email="b@example.com")
    _grant(db_session, user_a)
    _grant(db_session, user_b)

    r_a = _client(db_session, user_a).get(f"/orgs/{org_a}/exports/{AUTHORIZED_DEVICES.id}")
    r_b = _client(db_session, user_b).get(f"/orgs/{org_b}/exports/{AUTHORIZED_DEVICES.id}")
    assert r_a.status_code == 200 and r_b.status_code == 200
    cells_a, cells_b = _cell_values(r_a.content), _cell_values(r_b.content)
    assert "ALPHA-WS-01" in cells_a and "BRAVO-WS-01" not in cells_a
    assert "BRAVO-WS-01" in cells_b and "ALPHA-WS-01" not in cells_b


@pytest.mark.integration
def test_non_member_cannot_export_another_orgs_list(db_session):
    org_a = _seed_org_with_device(db_session, "ALPHA-WS-02")
    org_b = _seed_org_with_device(db_session, "BRAVO-WS-02")
    user_b = _make_fake_user(role="msp_engineer", org_id=org_b, email="b@example.com")
    _grant(db_session, user_b)

    r = _client(db_session, user_b).get(f"/orgs/{org_a}/exports/{AUTHORIZED_DEVICES.id}")
    assert r.status_code == 403


@pytest.mark.integration
def test_c3pao_assessor_can_export_every_view(db_session):
    org_id = _seed_org_with_device(db_session, "ASSESS-WS-01")
    assessor = _make_fake_user(role="c3pao_assessor", org_id=org_id)
    _grant(db_session, assessor)
    client = _client(db_session, assessor)

    for view in ALL_VIEWS:
        r = client.get(f"/orgs/{org_id}/exports/{view.id}")
        assert r.status_code == 200, f"{view.id}: {r.status_code} {r.text}"
        assert openpyxl.load_workbook(BytesIO(r.content)).active.title == view.sheet_title
    assert "ASSESS-WS-01" in _cell_values(
        client.get(f"/orgs/{org_id}/exports/{AUTHORIZED_DEVICES.id}").content
    )


@pytest.mark.integration
def test_c3pao_assessor_still_cannot_modify_a_scope_entity(db_session):
    org_id = _seed_org_with_device(db_session, "ASSESS-WS-02")
    assessor = _make_fake_user(role="c3pao_assessor", org_id=org_id)
    _grant(db_session, assessor)
    client = _client(db_session, assessor)
    entity_id = client.get(f"/orgs/{org_id}/scope").json()[0]["id"]

    assert client.patch(
        f"/orgs/{org_id}/scope/{entity_id}", json={"in_boundary": False}
    ).status_code == 403
    assert client.delete(f"/orgs/{org_id}/scope/{entity_id}").status_code == 403
    assert client.post(
        f"/orgs/{org_id}/scope",
        json={"entity_type": "device", "natural_key": "NEW-01", "attributes": {}},
    ).status_code == 403


def test_unknown_view_404s(monkeypatch):
    from fastapi import HTTPException

    with pytest.raises(HTTPException) as exc:
        scope_router.export_view(uuid.uuid4(), "not-a-view", session=None)
    assert exc.value.status_code == 404
    assert "not-a-view" not in VIEWS_BY_ID
